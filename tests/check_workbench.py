"""The workbench end to end on the synthetic pack (tests/make_pack.py), every model call stubbed:

  run       four files uploaded; the orchestrator reads the report's text, reads its key tables from their images (the
            discount rate is only in a pasted picture), extracts and reviews the key facts, checks each on the image of
            where it sits (the extraction swaps the rate's low and high, as a scrambled layout would; the image
            reading and the reviewer put it right),
            confirms the roles on the evidence (the rules' checks and a second opinion that agrees), rebuilds last
            year's overlay in Python (every formula cell as Excel saved it), finds the report's equity value (low
            and high on one row, ex-distribution), rolls forward onto this year's model, and writes the bridge
            (report -> rebuilt -> time value -> cash flows paid -> new forecast -> this year) and the cash-flow chart
            (last year's and this year's undiscounted forecast, the terminal value left out); gpt-sol reviews the run
  gating    a stage that fails isn't run again on the same inputs (only a person's "try again", or new inputs);
            gpt-sol decides an issue once per set of inputs, then a person does
  roles     where the second opinion disagrees, gpt-sol picks between the two on the evidence and the pick is
            checked in code before it's confirmed; a pick that fails the checks goes to a person
  inside    the same run with the overlay inside a copy of the client model and the report as slides
  workpaper the Excel workpaper built from the run, in memory: its sheets, the bridge's mid ends at this year's value,
            a row for every financial year of the cash flows, the disclaimer on its summary; none before the bridge
  held      inputs typed in the overlay outside its discountings: held at last year's, a suggestion from this year's
            model checked against last year's, applied when a person sets it, a bridge step of its own
  review    the review is a decision like the others: once per result, a try again doesn't buy another
  diagnostics  the run described in counts, ratios, dates and the app's own words: no name from the files
  rebuilt   this year's model rebuilt, with last year's figures pasted in under last year's labels: the copy is passed
            over, the rows found by their numbers, the value this year's; a point says the model looks rebuilt
  gate      this year's value held, with the reason, where it can't be trusted: a discounting left on last year's
            date, nothing of this year's model read; a big move at last year's date is a point to check
  lines     a later version of this year's model: a cash-flow line the overlay doesn't read (equity injections, on a
            sheet it reads) and the model saved on another scenario, each a point to check that leaves the value as
            it was; the ordinary run has neither (its scenario is last year's)
  changes   the finders behind those two, on small models: which lines count (not one a row the overlay reads adds
            up, one last year's model had figures in, or a rate) and which cells are selectors (not a note nothing
            reads, a series, or a case named in passing); a save time from the model's own cell, else the file's
  overview  every engagement at a glance: where it is (a finished one, an empty one), its values
  facts     the code checks behind the models, on the reviewers' cases: scale and currency, the label a figure sits
            under, dates and longer numbers, names, ranges, the image's waivers, table rows and headings, a blind read
  ranges    a dash between two figures is a range in the fact checks ("7.25% - 7.75%"), not a minus; a bracket
            or a dash after a word still is

    uv run python tests/make_pack.py && uv run python tests/check_workbench.py
"""
import json
import os
import re
import shutil
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "engine"))
PACK = ROOT / "tests" / "pack"

import calllog  # noqa: E402
import lessons  # noqa: E402
import library  # noqa: E402
import llm  # noqa: E402
import ratelimit  # noqa: E402
import usage  # noqa: E402
import workbench as wb  # noqa: E402
import workpaper  # noqa: E402
import orchestrator as orc  # noqa: E402


# ---- a sandbox: every file this run writes goes to a temp folder ----------------------------------------------------

def sandbox() -> Path:
    tmp = Path(tempfile.mkdtemp(prefix="workbench_"))
    wb.OUT, wb.UPLOADS, wb.DB, wb.DOCS = tmp / "out", tmp / "uploads", tmp / "out" / "workbench.db", tmp / "out" / "docs"
    library.OUT, library.UPLOADS, library.REGISTRY = tmp / "out", tmp / "uploads", tmp / "out" / "registry.db"
    usage.DB, calllog.DB, lessons.STORE = tmp / "out" / "registry.db", tmp / "out" / "calls.db", tmp / "out" / "lessons.json"
    (tmp / "out").mkdir(parents=True)
    ratelimit.load_capacities = lambda *a, **k: None
    return tmp


# ---- the models, stubbed -------------------------------------------------------------------------------------------

CALLS: list[str] = []
PROMPTS: dict[str, str] = {}  # the last prompt of each kind of call
BEHAVIOUR = {"second_opinion": "agree", "roles_decision": "rules", "verdict": "image", "verdict_figures": {}}
TABLES: dict[int, str] = {}  # page -> the table on it, as a reader of its image would transcribe it
TRUTH = {"equity_value": {"value": "A$2,507.9m", "low": "A$2,345.5m", "high": "A$2,670.3m"},
         "equity_value_cum": {"value": "A$2,532.9m"}, "valuation_date": {"value": "30 June 2025"},
         "discount_rate": {"low": "7.25%", "high": "7.75%"}, "terminal_growth_rate": {"value": "2.50%"},
         "franking_utilisation": {"value": "50%"}, "franking_credits_value": {"value": "A$389.8m"},
         "franking_credits_share": {"value": "15.5%"}, "terminal_value": {"value": "A$4,968.9m"},
         "pv_forecast": {"value": "A$1,819.3m"}, "pv_terminal_value": {"value": "A$1,173.7m"}}


class Reply:
    def __init__(self, out):
        self.output_text, self.usage = json.dumps(out), None


def _section(text: str, head: str, end: str | None) -> str:
    i = text.index(head) + len(head)
    return text[i:text.index(end, i)] if end else text[i:]


def _quote(page: str, needle: str) -> str:
    """The sentence (or table row, its cells separated by spaces) on the page holding the needle."""
    for line in page.splitlines():
        line = re.sub(r"\s+", " ", re.sub(r"\s*\|\s*", " ", line)).strip()
        if needle in line:
            for part in re.split(r"(?<=\.) (?=[A-Z])", line):
                if needle in part:
                    return part.strip()
    raise AssertionError(f"{needle!r} not on the page")


def fake_facts(doc: str) -> list[dict]:
    pages = {int(n): t for n, t in re.findall(r"<!-- page (\d+) -->\n(.*?)(?=<!-- page \d+ -->|\Z)", doc, re.S)}
    flat = lambda t: re.sub(r"\s+", " ", re.sub(r"\s*\|\s*", " ", t))
    where = lambda needle: next(n for n, t in pages.items() if needle in flat(t))

    def f(cat, key, label, needle, value_text="", low="", high="", unit="A$m", basis="", value=None):
        n = where(needle)
        return {"category": cat, "key": key, "label": label, "value_text": value_text, "low_text": low, "high_text": high,
                "value": value, "unit": unit, "basis": basis, "page": n, "quote": _quote(pages[n], needle)}
    return [f("identity", "target_name", "Target", "equity in Asset A Pty Ltd", "Asset A Pty Ltd", unit="text"),
            f("identity", "client", "Client", "engaged by Holdco A Pty Ltd", "Holdco A Pty Ltd", unit="text"),
            f("identity", "valuation_date", "Valuation date", "as at 30 June 2025 (the Valuation Date)", "30 June 2025",
              unit="date", value=20250630),
            f("conclusion", "equity_value", "Equity value", "A$2,345.5m", "A$2,507.9m", "A$2,345.5m", "A$2,670.3m",
              basis="ex-distribution"),
            f("conclusion", "equity_value_cum", "Equity value (cum-distribution)", "the midpoint would be A$2,532.9m",
              "A$2,532.9m", basis="cum-distribution"),
            # the low and the high the wrong way round, as a scrambled text layer can give them
            f("assumption", "discount_rate", "Discount rate", "Discount rate 7.25% - 7.75%", "", "7.75%", "7.25%",
              unit="%", basis="Post-tax nominal WACC"),
            f("assumption", "terminal_growth_rate", "Terminal growth rate", "Terminal growth rate 2.50%", "2.50%", unit="%"),
            f("assumption", "franking_utilisation", "Franking credit utilisation", "Franking credit utilisation 50%", "50%",
              unit="%", basis="Gamma"),
            f("conclusion", "franking_credits_value", "Value of franking credits", "franking credits of A$389.8m",
              "A$389.8m"),
            f("conclusion", "franking_credits_share", "Franking credits, share of equity value", "15.5% of the equity value",
              "15.5%", unit="%"),
            f("conclusion", "terminal_value", "Terminal value", "the terminal value is A$4,968.9m", "A$4,968.9m"),
            f("conclusion", "pv_forecast", "PV of the forecast", "A$1,819.3m", "A$1,819.3m"),
            f("conclusion", "pv_terminal_value", "PV of the terminal value", "A$1,173.7m", "A$1,173.7m"),
            f("assumption", "terminal_value_method", "Terminal value method", "using the Gordon growth method",
              "the Gordon growth method", unit="text")]


def fake_create(client, model, input, text=None, max_output_tokens=None, purpose=None, **kw):
    name = text["format"]["name"]
    CALLS.append(name)
    if isinstance(input, list):  # a call with an image: its prompt is the text part
        input = next(c["text"] for c in input[0]["content"] if c["type"] == "input_text")
    PROMPTS[name] = input
    if name == "table_read":
        page = int(re.search(r"(?:page|slide) (\d+)", input)[1])
        md = TABLES.get(page, "")
        return Reply({"is_table": bool(md), "title": "", "markdown": md, "description": "" if md else "a picture"})
    if name == "visual_read":
        items = json.loads(_section(input, "Items:\n", None))
        return Reply({"reads": [{"id": x["id"], "found": x["key"] in TRUTH, "column": "", "row": x["item"],
                                 **{k: TRUTH.get(x["key"], {}).get(k, "") for k in ("value", "low", "high")}} for x in items]})
    if name == "visual_verdict":
        items = json.loads(_section(input, "Items:\n", None))
        fig = BEHAVIOUR["verdict_figures"]
        return Reply({"verdicts": [{"id": x["id"], "choice": BEHAVIOUR["verdict"], "value": fig.get("value", ""),
                                    "low": fig.get("low", ""), "high": fig.get("high", ""),
                                    "reason": "the image shows 7.25% under Low and 7.75% under High"} for x in items]})
    if name == "identity":
        raise RuntimeError("no model in the test: identify falls back to the cells it found")
    if name == "report_facts":
        return Reply({"facts": fake_facts(_section(input, "Document:\n", None)), "notes": ""})
    if name == "facts_review":
        facts = json.loads(_section(input, "Facts:\n", "\n\nDocument:"))
        return Reply({"reviews": [{"id": f["id"], "verdict": "accept", "reason": "as the report states it", "value_text": "",
                                   "low_text": "", "high_text": "", "basis": "", "page": None, "quote": ""} for f in facts],
                      "missing": [], "summary": "all accepted"})
    if name == "roles_second_opinion":
        sug = json.loads(_section(input, "Suggested assignment:\n", None))
        if BEHAVIOUR["second_opinion"] == "agree":
            return Reply({"agree": True, **{k: sug.get(k, "") for k in ("prior_report", "prior_model", "prior_overlay",
                                                                         "current_model")},
                          "overlay_sheets": sug.get("overlay_sheets") or [], "confidence": "high", "reasons": ["as suggested"]})
        return Reply({"agree": False, **{k: sug.get(k, "") for k in ("prior_report", "prior_model", "prior_overlay")},
                      "current_model": "no such file.xlsx", "overlay_sheets": sug.get("overlay_sheets") or [],
                      "confidence": "medium", "reasons": ["a different current model"]})
    if name == "roles_decision":
        return Reply({"choice": BEHAVIOUR["roles_decision"], "reason": "the dates and links support it",
                      "question": "Which file is this year's client model?"})
    if name == "run_review":  # one point, about two years the chart has and one it doesn't, and a bridge step
        return Reply({"choice": "concerns", "reason": "one note", "question": "", "concerns": [
            {"title": "Terminal-year cash flow", "detail": "the new final year's cash flow jumps", "severity": "info",
             "years": ["FY2045", "FY2046", "FY2099"], "step": "forecast"}]})
    if name == "test_decision":
        return Reply({"choice": "pick", "reason": "test", "question": ""})
    if name == "row_action":  # the row agents, at their most careful: this year's model has no such row
        return Reply({"action": "not_in_this_model", "query": None, "row": None, "why": "no line item like it",
                      "confidence": "medium"})
    if name == "row_verdict":
        return Reply({"verdict": "reject", "why": "not the same line item", "better_row": None})
    if name == "row_advice":
        return Reply({"rows": [], "done": True})
    raise AssertionError(f"unexpected model call {name}")


def stub_models() -> None:
    llm.client = lambda *a, **k: None
    llm.create = fake_create
    import docingest  # what a reader of each page's table image sees: the slides' own tables, which are the same
    slides = docingest.process(str(PACK / "AssetA_valuation_report_FY25.pptx"), tempfile.mkdtemp(), read=False)
    TABLES.update({t["page"]: t["markdown"] for t in slides["tables"]})


# ---- helpers -------------------------------------------------------------------------------------------------------

def upload(eid: int, name: str, role: str | None = None) -> dict:
    src = PACK / name
    tmp = Path(tempfile.mkdtemp()) / name
    shutil.copy(src, tmp)
    return wb.add_upload(eid, tmp, name, library.sha256_file(tmp), role)


def wait(eid: int, done, what: str, timeout: float = 300) -> dict:
    t0 = time.time()
    while time.time() - t0 < timeout:
        v = orc.view(eid)
        if done(v):
            return v
        time.sleep(0.5)
    v = orc.view(eid)
    raise AssertionError(f"timed out waiting for {what}: " + json.dumps({s["stage"]: (s["status"], s["note"]) for s in v["stages"]},
                                                                        indent=1) + "\n" + "\n".join(
        f"{h['stage']}: {h['event']}: {h['text']}" for h in v["log"][:15]))


def status(v: dict) -> dict:
    return {s["stage"]: s["status"] for s in v["stages"]}


# ---- the checks ----------------------------------------------------------------------------------------------------

PACK_A = ("AssetA_valuation_report_FY25.pdf", "AssetA_BP25_client_model_Jun25.xlsx", "Alpha_valuation_overlay_FY25.xlsx",
          "AssetA_BP26_client_model_Jun26.xlsx")
PACK_R = ("AssetA_valuation_report_FY25.pdf", "AssetA_BP25_client_model_Jun25.xlsx", "Alpha_valuation_overlay_FY25.xlsx",
          "AssetA_FY26_plan_rebuilt.xlsx")  # this year's model rebuilt
PACK_B = ("AssetA_valuation_report_FY25.pptx", "AssetA_BP25_with_overlay.xlsx", "AssetA_BP26_client_model_Jun26.xlsx")
PACK_V = ("AssetA_valuation_report_FY25.pdf", "AssetA_BP25_client_model_Jun25.xlsx", "Alpha_valuation_overlay_FY25.xlsx",
          "AssetA_BP26_client_model_Jun26_v2.xlsx")  # a later version of this year's model


INSURANCE = "Operations!r10"  # this year's model adds insurance to EBITDA: a new term the value takes in


def confirm_insurance(eid: int) -> None:
    """A person's confirmation, ahead of the run, that this year's insurance belongs in the value (run_check sees
    the value held on it first)."""
    wb.confirm_term(eid, INSURANCE, True, "Insurance")


def run_check(files=PACK_A, name="Asset A, FY26") -> int:
    e = wb.create(name)
    eid = e["id"]
    n0 = dict((c, CALLS.count(c)) for c in set(CALLS))
    for f in files:
        upload(eid, f)
    v = wait(eid, lambda v: status(v).get("review") in (*orc.SETTLED, "blocked", "failed") or
             any(s["status"] in ("blocked", "failed") for s in v["stages"]), "the run", 600)
    st = status(v)
    # this year's model adds insurance to EBITDA, a term last year's didn't have: the value takes it in, so it waits
    # until a person confirms it belongs
    terms = [n for n in v["needs"] if n["id"] == "new-terms"]
    assert st["result"] == "blocked" and [n["rows"] for n in terms] == [[INSURANCE]], \
        (st, [(s["stage"], s["note"]) for s in v["stages"]], v["needs"])
    t0 = orc.stage(eid, "result").get("finished_at")
    confirm_insurance(eid)
    orc.poke()
    v = wait(eid, lambda v: orc.stage(eid, "result").get("finished_at") != t0 and not v["busy"] and
             status(v).get("review") in (*orc.SETTLED, "blocked", "failed"), "the run with the new term confirmed", 600)
    st = status(v)
    assert all(st[s] in orc.SETTLED for s in ("files", "facts", "roles", "rebuild", "rows", "result", "review")), \
        (st, [(s["stage"], s["note"]) for s in v["stages"]], v["needs"])
    e = wb.get(eid)
    rl = e["roles"]
    assert len(rl) == 4 and all(r["confirmed"] and r["by"] == "orchestrator" and r["evidence"] for r in rl.values()), rl
    names = {w["id"]: w["filename"] for w in e["workbooks"]}
    assert names[rl["prior_overlay"]["id"]] == files[2 if files == PACK_A else 1] and \
        names[rl["current_model"]["id"]].startswith("AssetA_BP26"), (names, rl)
    if files == PACK_B:
        assert rl["prior_model"]["id"] == rl["prior_overlay"]["id"] and \
            {"Val_Inputs", "DCF", "Summary"} <= set(rl["prior_overlay"]["sheets"]) and \
            "CashFlow" in rl["prior_model"]["sheets"], rl
    approved = {f["key"] for f in e["facts"] if f["status"] == "approved" and f["decided_by"] == "agents"}
    assert {"equity_value", "discount_rate", "valuation_date", "franking_utilisation"} <= approved, approved
    res = e["result"]
    assert res["head"]["basis"] == "ex" and abs(res["head"]["mid"] - 2507.9) < 1e-9, res["head"]
    assert res["tie"]["low"]["ok"] and res["tie"]["high"]["ok"] and res["tie"]["low"]["rebuilt_ok"], res["tie"]
    held = {h["cell"]: h["suggestion"]["status"] for h in res.get("held") or []}  # either layout: found, and checked
    assert held == {"Val_Inputs!C7": "checked", "Val_Inputs!C8": "checked"}, held
    # this year's model saved on last year's scenario, and no line in it the overlay doesn't read
    sc = res["scenario"]
    assert [(x["cell"], x["this_year"], x["last_cell"], x["last_year"], x["status"]) for x in sc["selectors"]] ==         [("Inputs!B14", 1.0, "Inputs!B13", 1.0, "same")] and sc["saved"]["this_year"]["date"] == "2026-08-12", sc
    assert not res["figures"]["gaps"]["new_lines"], res["figures"]["gaps"]["new_lines"]
    assert [(x["row"], x["label"], x["under"], x["in_value"], x["confirmed"], x["hold"]) for x in
            res["figures"]["gaps"]["new_terms"]] == [("Operations!r10", "Insurance", "Operations!r11", True, True, False)], \
        res["figures"]["gaps"]["new_terms"]
    assert not res["figures"]["gaps"]["gone_terms"], res["figures"]["gaps"]["gone_terms"]
    assert res["terminal"]["kind"] == "growth_final_year" and res["terminal"]["from"] == "the fact", res["terminal"]
    assert e["terminal"]["kind"] == "growth_final_year" and e["terminal"]["passages"], e["terminal"]
    assert not res["inputs"]["growth"].get("na") and res["inputs"]["growth"]["ok"], res["inputs"]["growth"]
    v = res["values"]
    assert abs(v["rebuilt"]["mid"] - 2507.876) < 0.01 and v["this_year"]["mid"] > v["rebuilt"]["mid"], v
    steps = res["bridges"]["mid"]["steps"]
    keys = [s["key"] for s in steps]
    assert keys == ["report", "rounding", "prior_feed", "rebuilt", "time", "cash", "forecast", "rate", "this_year"], keys
    moves = sum(s["value"] for s in steps if not s.get("total") and s["key"] not in ("rounding", "prior_feed"))
    assert abs(steps[3]["value"] + moves - steps[-1]["value"]) < 1e-6, "the steps add up"
    ch = res["chart"]
    assert ch["left_out"] == ["Terminal value"] and len(ch["series"]["last_year"]) == 20 and \
        len(ch["series"]["this_year"]) == 20, ch
    assert list(ch["series"]["last_year"])[0] == "FY2026" and list(ch["series"]["this_year"])[0] == "FY2027", ch["years"]
    assert ch["years"] == sorted(ch["years"]) and ch["years"][-1] == "FY2046", ch["years"]  # past 2049 sorts in order too
    # the models side by side: last year's report, overlay as saved and rebuild agree on every row
    sp = res["specs"]
    assert not sp.get("error") and [s["title"] for s in sp["sections"]][:3] == ["The value", "Assumptions", "Discounting"], sp
    rows = {r["label"]: r for s in sp["sections"] for r in s["rows"]}
    assert not [k for k, r in rows.items() if r["flag"]], [(k, r["cells"][:3]) for k, r in rows.items() if r["flag"]]
    assert rows["Equity value, mid"]["cells"][5]["v"] == res["values"]["this_year"]["mid"], rows["Equity value, mid"]
    assert rows["Undiscounted forecast cash flows (no terminal value)"]["cells"][3] and \
        rows["Undiscounted forecast cash flows (no terminal value)"]["cells"][4], rows["Undiscounted forecast cash flows (no terminal value)"]
    # what drives the value: the rate, the terminal value's growth, franking and the client's cash flows, traced
    dv = {g["group"]: g["inputs"] for g in res["drives"]["groups"]}
    assert not res["drives"].get("error") and {"Discount rate", "Terminal value", "Franking credits",
                                                "Cash flows (client model)"} <= set(dv), (res["drives"].get("error"), list(dv))
    assert all(x["effect"] is not None and x["effect"] < 0 for x in dv["Discount rate"]), dv["Discount rate"]
    assert any(x["effect"] and x["effect"] > 0 for x in dv["Franking credits"]), dv["Franking credits"]
    assert any(x["effect"] and x["effect"] > 0 and x.get("made_from") for x in dv["Cash flows (client model)"]), \
        dv["Cash flows (client model)"]
    # this year's cash flows against last year's: every discounting read period by period, the step split
    fl = res["flows"]
    assert fl["cores"] and all(c.get("last") and c.get("this") and c["form"]["recomputed"] for c in fl["cores"]), fl["cores"]
    sp = res["flow_checks"]["split"]
    assert set(sp) == {"low", "high"} and all(abs(s["outside"]) < 0.005 * 2507.9 for s in sp.values()), sp
    # the terminal value's part, on each feed, is the reconciliation's terminal value (the discrete rows fed too)
    low = next(c for c in fl["cores"] if "low" in c["ends"] and c.get("this"))
    for key, rk in (("last", "last_year"), ("this", "this_year")):
        tv = sum(p["tv"] for p in low[key]["periods"].values())
        assert abs(tv - res["reconcile"][rk]["low"]["tv"]) < 1e-6 * abs(tv), (key, tv, res["reconcile"][rk]["low"])
    disc = rows["Undiscounted forecast cash flows (no terminal value)"]["cells"]
    assert abs(disc[2]["v"] - disc[5]["v"]) > 1, disc  # this year's discrete forecast isn't last year's
    fc = [n for n in res["figures"]["gaps"]["holds"] if n["severity"] == "block"]
    assert not fc and res["flow_checks"]["rows"]["changed"] > 0, (fc, res["flow_checks"]["rows"])
    rate = next(f for f in e["facts"] if f["key"] == "discount_rate")
    if files == PACK_A:  # read from the picture, the swap put right on the image, and the picture's table read
        assert (rate["low_text"], rate["high_text"]) == ("7.25%", "7.75%") and rate["visual"]["status"] == "corrected", rate
        assert rate["status"] == "approved", rate
        seen = {f["key"]: (f.get("visual") or {}).get("status") for f in e["facts"]}
        assert seen["equity_value"] == "confirmed" and seen["terminal_value"] == "confirmed", seen
        doc = wb.document(rate["document_id"])["doc"]
        pic = next(t for t in doc["tables"] if t["source"] == "picture")
        assert pic["status"] == "verified" and "Discount rate" in pic["markdown"], pic
    else:
        assert rate["visual"]["status"] == "native", rate
    rec = {r["key"]: r for r in res["reconcile"]["rows"]}
    for key in ("terminal_value", "pv_forecast", "pv_terminal_value", "franking_credits_value", "franking_credits_share"):
        assert rec[key].get("ok") is True, (key, rec[key])
    this = res["reconcile"]["this_year"]["mid"]
    assert this["pv_tv"] and this["pv_forecast"] and abs(this["pv"] - (this["pv_tv"] + this["pv_forecast"])) < 1e-6, this
    rates = {e: next(a["rate_source"] for a in res["assumptions"][e] if a.get("parts")) for e in ("low", "high")}
    assert rates == {"low": "Val_Inputs!C5", "high": "Val_Inputs!E5"}, rates
    # each end's rate sighted, sourced by following the factors' formulas, and checked again against the report
    ins = res["inputs"]
    rc = ins["rate"]["ends"]
    assert ins["rate"]["ok"] and [(rc[e]["cell"], rc[e]["sourced"], rc[e]["ties"], rc[e]["report"], rc[e]["sighted"]["heading"],
                                   rc[e]["sighted"]["input"]) for e in ("low", "high")] == \
        [("Val_Inputs!C5", True, True, "7.75%", "Low", True), ("Val_Inputs!E5", True, True, "7.25%", "High", True)], rc
    # the terminal growth rate: the cell the terminal value reads as g, the terminal value at each end's own rate
    gr = ins["growth"]["ends"]
    assert ins["growth"]["ok"] and all(gr[e]["cell"] == "Val_Inputs!C6" and gr[e]["ties"] and
                                       all(c["ok"] for c in gr[e]["checks"]) for e in ("low", "high")), gr
    # franking credit utilisation: the fraction the franking credits read; at nil, the equity value falls by exactly
    # the value of franking credits in the Python overlay, and is back where it was after
    fk = ins["franking"]["ends"]
    assert ins["franking"]["ok"] and all(fk[e]["cell"] == "Val_Inputs!C9" and fk[e]["ties"] and fk[e]["rerun"]["ok"] and
                                         fk[e]["rerun"]["restored"] for e in ("low", "high")), fk
    assert abs(fk["low"]["rerun"]["drop"] - 381.1203152392057) < 1e-6, fk["low"]["rerun"]
    # the date line: the report's date the anchor, the overlay's the cell its discount factors read
    dl = {d["step"]: d for d in wb.get(eid)["dates"]}
    assert (dl["The report"]["date"], dl["Last year's overlay"]["date"], dl["Last year's overlay"]["where"],
            dl["Last year's overlay"]["ok"], dl["Last year's client model"]["ok"], dl["This year's client model"]["date"],
            dl["This year's client model"]["ok"], dl["The roll-forward runs to"]["date"]) == \
        ("2025-06-30", "2025-06-30", "Val_Inputs!C4", True, True, "2026-06-30", True, "2026-06-30"), dl
    view = orc.view(eid)
    assert not [n for n in view["needs"] if n["severity"] == "block" or n["id"] in ("new-lines", "scenario")], view["needs"]
    # every need says what it asks of a person and lands on a card; a review point names its years and step, checked
    # against the run (a year the chart doesn't have is dropped)
    assert all(n.get("kind") and (n["kind"] == "retry" or n["go"].get("anchor")) for n in view["needs"]), view["needs"]
    rp = next(n for n in view["needs"] if n["id"] == "review-0")
    assert (rp["kind"], rp["years"], rp["step"], rp["go"]["anchor"]) == ("review-point", ["FY2045", "FY2046"], "forecast", "review-0"), rp
    # how long each stage took, so the page can say how long is left
    took = {x["stage"]: x["typical_secs"] for x in view["stages"]}
    assert all(took[k] is not None for k in ("facts", "roles", "rebuild", "rows", "result", "review")), took
    log = orc.history(eid, limit=200)
    assert any(h["stage"] == "roles" and h["event"] == "done" for h in log) and \
        any(h["stage"] == "review" and h["event"] == "decide" for h in log)
    calls = {c: CALLS.count(c) - n0.get(c, 0) for c in set(CALLS)}
    assert calls.get("report_facts") == 1 and calls.get("facts_review") == 1 and not calls.get("roles_decision"), calls
    if files == PACK_A:
        assert calls.get("visual_read") == 3 and calls.get("visual_verdict") == 1, calls
        # what the report says around the figures reached every model that read or judged one
        for kind in ("visual_read", "visual_verdict", "table_read"):
            assert "Dear Directors" in PROMPTS[kind] and "context only" in PROMPTS[kind], kind
        assert "Scope of our work" in PROMPTS["visual_verdict"] and "Established so far" in PROMPTS["visual_verdict"]
    print(f"{'run' if files == PACK_A else 'inside'}: ok (the terminal value, the PV split and the franking credits "
          f"reconcile to the report; facts agreed, roles confirmed on {len(rl['current_model']['evidence']['checks'])} checks and a "
          f"second opinion; {res['where']['low']} / {res['where']['high']} tie; mid {v['report']['mid']:,.1f} -> "
          f"{v['this_year']['mid']:,.1f}; chart {len(ch['series']['last_year'])} + {len(ch['series']['this_year'])} years; "
          f"model calls {calls})")
    return eid


def gating_check(eid: int) -> None:
    runs = []

    def broken(eid_, key):
        runs.append(key)
        raise RuntimeError("broken on purpose")
    keep = orc.STAGE_JOBS["map"]
    orc.STAGE_JOBS["map"] = broken
    try:
        orc.retry(eid, "map")
        wait(eid, lambda v: status(v)["map"] == "failed", "the map to fail")
        for _ in range(4):
            orc.tick(eid)
            time.sleep(0.3)
        assert len(runs) == 1, f"a failed stage ran again on the same inputs: {len(runs)}"
        assert any(n["id"] == "failed-map" for n in orc.view(eid)["needs"])
        orc.retry(eid, "map")
        wait(eid, lambda v: len(runs) == 2 and status(v)["map"] == "failed", "the person's try again")
    finally:
        orc.STAGE_JOBS["map"] = keep
    # the roles run again and come out the same: the rebuild and what follows keep their outcome, without running again
    before = {n: orc.stage(eid, n)["started_at"] for n in ("rebuild", "rows", "result")}
    orc.retry(eid, "roles")
    wait(eid, lambda v: all(status(v)[n] in orc.SETTLED for n in ("roles", "rebuild", "rows", "result")), "the rerun")
    after = {n: orc.stage(eid, n)["started_at"] for n in ("rebuild", "rows", "result")}
    assert before == after, (before, after)
    schema = orc._schema("test_decision", ["pick", "escalate"])
    n0 = CALLS.count("test_decision")
    a = orc.ask(eid, "map", "test", "k1", "decide", schema)
    b = orc.ask(eid, "map", "test", "k1", "decide", schema)
    c = orc.ask(eid, "map", "test", "k2", "decide", schema)
    assert a and b is None and c and CALLS.count("test_decision") == n0 + 2, (a, b, c)
    assert any(h["event"] == "escalate" and h["issue"] == "test" for h in orc.history(eid, "map"))
    orc.retry(eid, "map")
    wait(eid, lambda v: status(v)["map"] == "done", "the map")
    print("gating: ok (a failed stage waits for new inputs or a person's try again; a stage the rerun before it didn't "
          "change keeps its outcome; gpt-sol decides an issue once per set of inputs, then escalates)")


def roles_check() -> None:
    BEHAVIOUR.update(second_opinion="differ", roles_decision="rules")
    e = wb.create("Asset A, FY26 (roles)")
    eid = e["id"]
    for name in ("AssetA_valuation_report_FY25.pdf", "AssetA_BP25_client_model_Jun25.xlsx",
                 "Alpha_valuation_overlay_FY25.xlsx", "AssetA_BP26_client_model_Jun26.xlsx"):
        upload(eid, name)
    wait(eid, lambda v: status(v)["roles"] in (*orc.SETTLED, "blocked", "failed"), "the roles")
    rl = wb.roles(eid)
    assert all(r["by"] == "orchestrator" and "decided_by" in (r["evidence"] or {}) for r in rl.values()), rl
    assert any(h["event"] == "verify" for h in orc.history(eid, "roles"))
    # gpt-sol takes the second opinion, which names a file that isn't here: the checks in code stop it
    BEHAVIOUR.update(roles_decision="second_opinion")
    with wb._lock, wb._conn() as db:
        db.execute("DELETE FROM roles WHERE engagement_id=?", (eid,))
    orc.retry(eid, "roles")
    v = wait(eid, lambda v: status(v)["roles"] == "blocked", "the roles to go to a person")
    need = next(n for n in v["needs"] if n["id"] == "roles")
    assert need["severity"] == "block" and "Which file" in need["detail"], need
    BEHAVIOUR.update(second_opinion="agree", roles_decision="rules")
    print("roles: ok (a disagreement is decided by gpt-sol on the evidence and checked in code; a pick that fails the "
          "checks goes to a person)")


def adviser_tabs_check() -> None:
    """An overlay in a copy of the client model, its sheets behind the adviser's divider tab ("Adviser>>", up to "Client>>"):
    those are its own, the rest its copy of the client model, where last year's client model has the same sheets (even
    with the adviser's tabs in it too); a standalone overlay's own dividers ("Outputs >") leave every sheet its own."""
    import build_map
    import roles as rolesmod
    import xlsxwriter
    out = Path(tempfile.mkdtemp(prefix="tabs_"))
    book = xlsxwriter.Workbook(out / "copy.xlsx")
    for name, rows in (("Adviser>>", 0), ("Valuation", 4), ("PFI", 3), ("Client>>", 0), ("Inputs", 3), ("Ops", 5),
                       ("Fin", 4), ("SPARE>>", 0), ("Scratch", 2)):
        ws = book.add_worksheet(name)
        ws.write(0, 0, name.strip(">") + " section" if not rows else name)
        for r in range(rows):
            ws.write(r + 2, 1, f"{name} line {r}")
            ws.write_number(r + 2, 3, 10.0 + r)
    book.close()
    db = build_map.main(str(out / "copy.xlsx"), str(out / "db"))["db"]
    names = ["Adviser>>", "Valuation", "PFI", "Client>>", "Inputs", "Ops", "Fin", "SPARE>>", "Scratch"]
    assert rolesmod.adviser_sheets(db, names, names) == ["Adviser>>", "Valuation", "PFI"]
    assert rolesmod.adviser_sheets(db, names, ["Inputs", "Ops", "Fin", "Scratch"]) == ["Adviser>>", "Valuation", "PFI"]
    # a standalone overlay: little of the rest is the client model's
    assert rolesmod.adviser_sheets(db, names, ["Ops"]) is None
    # a firm's own name for its divider is a setting (brand/settings.json, here a temporary one), not in the code
    import likeness
    saved, env = likeness.SETTINGS, os.environ.pop("VALUATION_DESK_OVERLAY_MARKERS", None)
    likeness.SETTINGS = out / "settings.json"
    try:
        named = ["XY>>" if s == "Adviser>>" else s for s in names]
        assert rolesmod.adviser_sheets(db, named, names, rolesmod.adviser_name()) is None
        assert likeness.save_markers([" XY ", "xy", ""]) == ["XY"] and likeness.markers() == ["xy"]
        assert rolesmod.adviser_sheets(db, named, names, rolesmod.adviser_name()) == ["XY>>", "Valuation", "PFI"]
        assert not rolesmod.adviser_name().match("XYZ") and not rolesmod.adviser_name().match("Valuation")
    finally:
        likeness.SETTINGS = saved
        if env is not None:
            os.environ["VALUATION_DESK_OVERLAY_MARKERS"] = env
    print("adviser tabs: ok (the sheets behind the Adviser>> tab are the overlay's own, the rest its copy of the client "
          "model; a standalone overlay's tabs leave it whole; a firm's own name for the tab is a setting kept on this "
          "machine)")


def copies_check() -> None:
    """Two copies of last year's client model with the adviser's tabs behind "Adviser>>" (the one the report came from and
    an earlier working copy) and this year's client model, its sheets renamed and rebuilt: unlike the others, with
    valuation words and charts, and as many of the report's figures (last year's, kept in its history). The overlay is
    the report's copy's tabs, though the earlier copy has the same sheets, the tabs tipping it; last year's client model
    is that copy's own client sheets, not the earlier copy. The tabs are one sign among others, not every overlay being
    laid out so: a working copy with them doesn't outweigh a separate valuation workbook holding the report's figures,
    and a client model's own "Valuation>" section isn't the adviser's."""
    import build_map
    import roles as rolesmod
    import xlsxwriter
    from datetime import datetime
    from xlsxwriter.utility import xl_rowcol_to_cell as cell
    out = Path(tempfile.mkdtemp(prefix="copies_"))

    def book(i, name, sheets, vd, charts=()):
        b = xlsxwriter.Workbook(out / name)
        day = b.add_format({"num_format": "dd mmm yyyy"})
        for s, rows in sheets:
            ws = b.add_worksheet(s)
            ws.write(0, 0, s.strip("<> ") + " section" if not rows else s)
            for r, (label, v) in enumerate(rows, 2):
                ws.write(r, 1, label)
                if isinstance(v, datetime):
                    ws.write_datetime(r, 3, v, day)
                elif isinstance(v, tuple):  # six periods, each grown from the one before
                    ws.write_number(r, 3, v[0])
                    for c in range(4, 9):
                        ws.write_formula(r, c, f"={cell(r, c - 1)}*{v[1]}")
                else:
                    ws.write(r, 3, v)
            if s in charts:
                ch = b.add_chart({"type": "line"})
                ch.add_series({"values": [s, 2, 3, 2, 8]})
                ws.insert_chart("K2", ch)
        b.close()
        db = build_map.main(str(out / name), str(out / f"db{i}"))["db"]
        return {"id": i, "filename": name, "db_path": db, "source_path": str(out / name), "valuation_date": vd,
                "uploaded_at": float(i)}

    model = [("Inputs", [("Volume growth", 0.025), ("Unit charge", 20.0), ("Tax rate", 0.3)]),
             ("Ops", [("Volumes", (10.0, 1.025)), ("Regulated revenue", (200.0, 1.03)),
                      ("Operating costs", (60.0, 1.025)), ("Present value of terminal value", 1900.0)]),
             ("CashFlow", [("EBITDA", (140.0, 1.03)), ("Capital expenditure", (50.0, 1.01)),
                           ("Distributions", (80.0, 1.04))]),
             ("Hist", [("Discount rate", 0.10), ("Historical revenue", (190.0, 1.04))]),
             ("Checks", [("Balance check", "=CashFlow!D3-CashFlow!D3")])]
    client = [("Client>>", [])] + model
    adv = lambda equity, growth, tv: [
        ("Adviser>>", []),
        ("Valuation", [("Valuation date", datetime(2025, 6, 30)), ("Distributions to equity", "=CashFlow!D5"),
                       ("Discount factor", (0.95, 0.91)), ("Terminal growth rate", growth), ("Terminal value", tv),
                       ("Equity value", equity)]),
        ("PFI", [("Revenue", "=Ops!D4"), ("EBITDA", "=CashFlow!D3")]),
        ("Adviser Charts", [("Equity value low", equity * 0.95), ("Equity value high", equity * 1.05)]),
        ("Sensitivities", [("Equity value at a 0.5% higher rate", equity * 0.93)])]
    current = [("Title", [("Asset A valuation model", "December 2025")]),
               ("Outputs>", []),
               ("Dashboard_O", [("Enterprise value", "=aVal!D3"), ("Equity value", "=aVal!D4")]),
               ("Ops_O", [("Volume", (10.5, 1.025)), ("Regulated revenue", (210.0, 1.03)),
                          ("Operating costs", (62.0, 1.025))]),
               ("Inputs>", []),
               ("iOps", [("Volume growth assumption", 0.025), ("Unit charge", 21.0)]),
               ("Analysis>", []),
               ("aVal", [("Enterprise value", 3000.0), ("Equity value", 1500.0), ("Net debt", "=D3-D4"),
                         ("Terminal year", 2050), ("Discount factor", (0.95, 0.91)), ("Gearing", "=D5/D3")]),
               ("System>", []),
               ("Hist", [("Discount rate", 0.10), ("Terminal growth rate", 0.03)]),
               ("Log", [("Valuation date", datetime(2025, 6, 30)), ("Model updated", datetime(2025, 12, 31))])]
    # this year's model first: on the report's figures alone it ties with the report's copy
    wbs = [book(3, "Asset A Valuation Model December 2025.xlsx", current, "2025-12-31", charts=("Dashboard_O",)),
           book(1, "Asset A June 2025 adviser overlay v2.xlsx", adv(1500.0, 0.03, 20000.0) + client, "2025-06-30"),
           book(2, "20250529 Asset A June 25 Valuation.xlsx", adv(1400.0, 0.025, 19000.0) + client, "2025-06-30")]
    facts = [{"id": 1, "category": "identity", "key": "valuation_date", "label": "Valuation date",
              "value_text": "30 June 2025", "value": 20250630, "unit": "date"},
             {"id": 2, "category": "conclusion", "key": "equity_value", "label": "Equity value", "value_text": "1,500.0",
              "unit": "A$m"},
             {"id": 3, "category": "conclusion", "key": "terminal_value", "label": "Terminal value", "value_text": "20,000",
              "unit": "A$m"},
             {"id": 4, "category": "conclusion", "key": "pv_terminal_value", "label": "Present value of terminal value",
              "value_text": "1,900", "unit": "A$m"},
             {"id": 5, "category": "assumption", "key": "discount_rate", "label": "Discount rate", "value_text": "10.0%",
              "unit": "%"},
             {"id": 6, "category": "assumption", "key": "terminal_growth_rate", "label": "Terminal growth rate",
              "value_text": "3.0%", "unit": "%"}]
    res = rolesmod.suggest([{"id": 1, "filename": "Asset A valuation report June 2025.pdf", "n_facts": 6}], wbs, facts)
    rl, by = res["roles"], res["workbooks"]
    assert by[3]["structure"][0].startswith("unlike the other workbooks"), by[3]["structure"]
    assert (rl["prior_overlay"]["id"], rl["prior_overlay"]["sheets"]) == \
        (1, ["Adviser>>", "Valuation", "PFI", "Adviser Charts", "Sensitivities"]), rl["prior_overlay"]
    assert (rl["prior_model"]["id"], rl["prior_model"]["sheets"]) == \
        (1, ["Inputs", "Ops", "CashFlow", "Hist", "Checks"]), rl["prior_model"]
    assert any("20250529 Asset A June 25 Valuation.xlsx has the adviser's tabs too" in w for w in rl["prior_model"]["why"])
    assert rl["current_model"]["id"] == 3 and by[3]["mode"] == "client model", (rl.get("current_model"), by[3])
    # the tabs are one sign: the client's models with a valuation section of their own, an early working copy with
    # the adviser's tabs, and the separate valuation workbook that holds the report's figures
    own = [("Valuation>", []), ("Val_Calc", [("Discount rate", 0.10), ("Equity IRR", 0.112), ("Project NPV", (55.0, 1.02))])]
    standalone = [("Val_Inputs", [("Valuation date", datetime(2025, 6, 30)), ("Discount rate", 0.10),
                                  ("Terminal growth rate", 0.03)]),
                  ("DCF", [("Distributions", (80.0, 1.04)), ("Discount factor", (0.95, 0.91)), ("Terminal value", 20000.0),
                           ("Present value of terminal value", 1900.0)]),
                  ("Summary", [("Enterprise value", 2350.0), ("Net debt", 850.0), ("Equity value", 1500.0),
                               ("Sensitivity to the discount rate", 120.0)])]
    working = [("Adviser>>", []), ("Adviser Notes", [("Valuation date", datetime(2025, 6, 30)), ("Equity value draft", 1250.0)]),
               ("Client>>", [])]
    wbs = [book(11, "Asset B client model June 2025.xlsx", model + own, "2025-06-30"),
           book(12, "Asset B client model June 2026.xlsx", model + own, "2026-06-30"),
           book(13, "Asset B model adviser working copy.xlsx", working + model + own, "2025-06-30"),
           book(14, "Asset B valuation FY25.xlsx", standalone, "2025-06-30")]
    res = rolesmod.suggest([{"id": 1, "filename": "Asset B valuation report June 2025.pdf", "n_facts": 6}], wbs, facts)
    rl, by = res["roles"], res["workbooks"]
    assert by[13]["tabs"] == ["Adviser>>", "Adviser Notes"] and by[11]["tabs"] is None and by[12]["tabs"] is None, by
    assert (rl["prior_overlay"]["id"], rl["prior_overlay"]["sheets"]) == (14, ["Val_Inputs", "DCF", "Summary"]), rl["prior_overlay"]
    assert rl["prior_model"]["id"] == 11 and rl["current_model"]["id"] == 12, (rl.get("prior_model"), rl.get("current_model"))
    print("copies: ok (two copies with the adviser's tabs: the overlay is the tabs of the one holding the report's "
          "figures, last year's client model its own client sheets, this year's renamed model a client model; the tabs "
          "one sign among others: a working copy with them doesn't outweigh the valuation workbook holding the figures, "
          "a client's own \"Valuation>\" section isn't the adviser's)")


def escalate_check() -> None:
    """Where the image reading and the reviewer agree with neither each other nor the extraction, the fact goes to a
    person: not approved, the agents' status escalated, on the needs-you list with the image."""
    BEHAVIOUR.update(verdict="neither", verdict_figures={"low": "9.00%", "high": "9.50%"})
    try:
        e = wb.create("Asset A, FY26 (the image disagrees)")
        eid = e["id"]
        for f in PACK_A:
            upload(eid, f)
        v = wait(eid, lambda v: status(v)["facts"] in (*orc.SETTLED, "failed"), "the facts")
        rate = next(f for f in wb.facts(eid) if f["key"] == "discount_rate")
        assert rate["status"] == "pending" and rate["agent"]["status"] == "escalated" and \
            rate["visual"]["status"] == "escalated" and rate["agent"]["open"]["image"].startswith("tables/"), rate
        need = next(n for n in v["needs"] if n["id"] == f"fact-{rate['id']}")
        assert need["severity"] == "check" and "on the image" in need["detail"], need
        for f in wb.facts(eid):  # nothing approved yet: the facts that pass their checks lead, but not one a person owes
            wb._set("facts", f["id"], status="pending")
        assert rate["id"] not in [f["id"] for f in wb.reference(eid)] and wb.reference(eid), wb.reference(eid)
    finally:
        BEHAVIOUR.update(verdict="image", verdict_figures={})
    print("escalate: ok (where the reads of the image agree with nothing, the fact goes to a person, with the image; "
          "it doesn't lead meanwhile)")


def place_check() -> None:
    """Files placed in their roles by hand while uploading stay the person's; the orchestrator places the rest (last
    year's client model here) and confirms it on the evidence. A file of the wrong kind for its role is refused before
    anything is stored."""
    e = wb.create("Asset A, FY26 (placed by hand)")
    eid = e["id"]
    bad = Path(tempfile.mkdtemp()) / PACK_A[0]
    shutil.copy(PACK / PACK_A[0], bad)
    try:
        wb.add_upload(eid, bad, PACK_A[0], library.sha256_file(bad), "current_model")
        raise AssertionError("a report was taken as this year's client model")
    except ValueError as ex:
        assert "a model goes in as XLSX or XLSM" in str(ex) and not bad.exists() and not wb.documents(eid), ex
    for name, role in zip(PACK_A, ("prior_report", None, "prior_overlay", "current_model")):
        upload(eid, name, role)
    wait(eid, lambda v: status(v)["roles"] in (*orc.SETTLED, "blocked", "failed"), "the roles")
    rl = wb.roles(eid)
    assert {k: (r["confirmed"], r["by"]) for k, r in rl.items()} == {
        "prior_report": (True, "you"), "prior_overlay": (True, "you"), "current_model": (True, "you"),
        "prior_model": (True, "orchestrator")}, rl
    print("place: ok (files placed by hand while uploading stay yours; the orchestrator places the rest; a file of the "
          "wrong kind for its role is refused before it's stored)")


def overview_check(eid: int) -> None:
    """The list of every engagement: where each is, and the value the result page shows."""
    e = wb.create("Asset A, nothing yet")
    try:
        ov = {x["id"]: x for x in orc.overview()}
        x, empty = ov[eid], ov[e["id"]]
        want = wb.get(eid)["result"]["values"]
        assert x["state"] == "finished" and x["values"] == want and x["units"] and x["basis"] == "ex" and \
            x["valuation_date"], x
        assert x["needs"] == len([n for n in wb.get(eid)["run"]["needs"] if n["severity"] != "info"]), x
        assert empty["state"] == "empty" and empty["values"] is None and empty["files"] == 0, empty
    finally:
        wb.delete(e["id"])
    print("overview: ok (a finished engagement with its values and what's for a person; an empty one as empty)")


def rows_context_check(eid: int) -> None:
    """A row to find comes with what it is: the heading it sits under and its neighbours, what its formula works out
    from, last year's figures, which overlay row reads it and whether that's a cash flow the value discounts; and
    this year's candidates with their figures for the same periods."""
    info = wb.row_info(eid, "CashFlow!r9")
    c = info["context"]
    assert c["heading"] == "Cash flow" and c["above"]["label"] == "Tax paid", c
    assert [x["label"] for x in c["formula"]["reads"]] == ["EBITDA", "Capital expenditure", "Tax paid"], c["formula"]
    assert len(c["values"]) == 5 and c["values"][0]["period"] == "2026-06-30", c["values"]
    assert c["read_by"] and c["read_by"][0]["row"] == "DCF!r5" and c["read_by"][0]["cells"] == 20 and c["feeds_dcf"], c
    k = info["candidates"][0]
    assert k["row"] == "CashFlow!r9" and k["figures"][0] is None and isinstance(k["figures"][1], float), k
    print("rows: ok (a row to find says what it is: its heading and neighbours, what it's worked out from, last "
          "year's figures, the overlay row that reads it; this year's candidates with their figures)")


def held_check(eid: int) -> None:
    """The inputs typed in the overlay outside its discountings (net debt, a declared distribution) are listed as held
    at last year's, each with a suggestion from this year's client model found through the row that holds last year's
    figure in last year's model (checked), and a point to check. Nothing changes until a person sets this year's
    figure: then the bridge has a step of its own for it, and this year's value moves by exactly that; back to last
    year's, it's the first result again."""
    import make_pack
    res = wb.get(eid)["result"]
    by = {h["cell"]: h for h in res["held"]}
    nd, dist = by["Val_Inputs!C7"], by["Val_Inputs!C8"]
    assert (nd["value"], dist["value"]) == (make_pack.NET_DEBT, make_pack.DISTRIBUTION) and nd["held"] and dist["held"], by
    assert nd["suggestion"]["status"] == dist["suggestion"]["status"] == "checked", (nd["suggestion"], dist["suggestion"])
    assert (nd["suggestion"]["value"], dist["suggestion"]["value"]) == (make_pack.NET_DEBT_NEW, make_pack.DISTRIBUTION_NEW)
    assert nd["suggestion"]["last"]["value"] == make_pack.NET_DEBT and "Less: net debt" in [x["label"] for x in nd["lines"]]
    needs = [n for n in orc.view(eid)["needs"] if n["id"].startswith("held-")]
    assert len(needs) == 2 and all(n["severity"] == "check" and n["go"]["anchor"] == "heldCard" for n in needs), needs
    assert "held" not in [x["key"] for x in res["bridges"]["mid"]["steps"]]
    mid0 = res["values"]["this_year"]["mid"]
    finished = lambda: orc.stage(eid, "result").get("finished_at")
    t0 = finished()
    wb.set_held(eid, "Val_Inputs!C7", nd["suggestion"]["value"], "suggestion")
    wb.set_held(eid, "Val_Inputs!C8", 30.0, "typed")
    orc.poke()
    wait(eid, lambda v: finished() != t0 and not v["busy"] and status(v)["result"] in orc.SETTLED, "the result with this year's inputs")
    res = wb.get(eid)["result"]
    move = -(make_pack.NET_DEBT_NEW - make_pack.NET_DEBT) - (30.0 - make_pack.DISTRIBUTION)
    for end in ("low", "mid", "high"):
        step = next(x for x in res["bridges"][end]["steps"] if x["key"] == "held")
        assert abs(step["value"] - move) < 1e-6, (end, step)
    assert abs(res["values"]["this_year"]["mid"] - (mid0 + move)) < 1e-6, (res["values"]["this_year"], mid0, move)
    assert not [n for n in orc.view(eid)["needs"] if n["id"].startswith("held-")]
    import io
    import openpyxl
    book = openpyxl.load_workbook(io.BytesIO(workpaper.build(eid)), data_only=True)
    rows = [r for r in book["Inputs"].iter_rows(values_only=True) if r[0] == "Net debt at valuation date"]
    assert rows and rows[0][2] == make_pack.NET_DEBT_NEW and "accepted from the model" in rows[0][3], rows
    assert any(r[6] == "Held inputs" for r in book["Bridge"].iter_rows(values_only=True)), "no held step in the chart"
    assert {h["cell"]: (h["held"], h["this_year"], h["from"]) for h in res["held"]} == \
        {"Val_Inputs!C7": (False, make_pack.NET_DEBT_NEW, "suggestion"), "Val_Inputs!C8": (False, 30.0, "typed")}
    t0 = finished()
    for c in ("Val_Inputs!C7", "Val_Inputs!C8"):
        wb.set_held(eid, c, None)
    orc.poke()
    wait(eid, lambda v: finished() != t0 and not v["busy"] and status(v)["review"] in orc.SETTLED, "the result back again")
    res = wb.get(eid)["result"]
    assert abs(res["values"]["this_year"]["mid"] - mid0) < 1e-9 and all(h["held"] for h in res["held"])
    print(f"held: ok (net debt and the declared distribution held at last year's, each with a suggestion from this year's "
          f"model checked against last year's; set, the bridge has their step ({move:+.1f}) and this year's value moves "
          f"by it; back, the first result)")


def rate_check(eid: int) -> None:
    """This year's discount rate, set by a person: the low end at the higher rate whichever order it's typed in, on the
    cells each end's discountings read. The roll-forward's steps stay at last year's rate and the rate is a step of its
    own; the zero-roll check stays at last year's rate (a new rate isn't a row matched wrongly); back to last year's,
    it's the first result again."""
    res = wb.get(eid)["result"]
    ends0 = res["inputs"]["rate"]["ends"]
    was = {e: ends0[e]["value"] for e in ("low", "high")}
    steps0 = {e: {x["key"]: x["value"] for x in res["bridges"][e]["steps"]} for e in ("low", "mid", "high")}
    zero0 = {c: x["zero_roll"]["ratio"] for c, x in res["figures"]["gaps"]["by_cell"].items()}
    assert all(abs(steps0[e]["rate"]) < 1e-12 for e in steps0), steps0
    mid0 = res["values"]["this_year"]["mid"]
    new = {"low": was["low"] + 0.0025, "high": was["high"] + 0.0025}
    finished = lambda: orc.stage(eid, "result").get("finished_at")
    t0 = finished()
    got = wb.set_this_year_rate(eid, f"{100 * new['high']:.2f}%", 100 * new["low"])  # the lower rate first, as typed
    assert abs(got["low"] - new["low"]) < 1e-12 and abs(got["high"] - new["high"]) < 1e-12, got
    orc.poke()
    wait(eid, lambda v: finished() != t0 and not v["busy"] and status(v)["result"] in orc.SETTLED, "the result at this year's rate")
    res = wb.get(eid)["result"]
    rt = res["inputs"]["rate"]["this_year"]
    assert rt["applied"] and rt["cells"] == {ends0["low"]["cell"]: got["low"], ends0["high"]["cell"]: got["high"]} \
        or (ends0["low"]["cell"] == ends0["high"]["cell"]), rt
    for e in ("low", "mid", "high"):
        s = {x["key"]: x for x in res["bridges"][e]["steps"]}
        for k in ("time", "cash", "forecast"):  # at last year's rate, as before
            assert abs(s[k]["value"] - steps0[e][k]) < 1e-6, (e, k, s[k]["value"], steps0[e][k])
        assert s["rate"]["value"] < 0 and "this year's, " in s["rate"]["label"] and "last year's" in s["rate"]["label"], s["rate"]
        assert abs(s["this_year"]["value"] - (steps0[e]["this_year"] + s["rate"]["value"])) < 1e-6, (e, s)
    assert {c: x["zero_roll"]["ratio"] for c, x in res["figures"]["gaps"]["by_cell"].items()} == zero0
    assert res["values"]["this_year"]["mid"] < mid0 and not [n for n in orc.view(eid)["needs"] if n["id"] == "rate-this-year"]
    import io
    import openpyxl
    book = openpyxl.load_workbook(io.BytesIO(workpaper.build(eid)), data_only=True)
    assert any(r[0] == "This year" and "set by you" in str(r[2]) for r in book["Inputs"].iter_rows(values_only=True))
    move = res["bridges"]["mid"]["steps"][-2]["value"]
    t0 = finished()
    wb.set_this_year_rate(eid, None)
    orc.poke()
    wait(eid, lambda v: finished() != t0 and not v["busy"] and status(v)["review"] in orc.SETTLED, "the result back again")
    res = wb.get(eid)["result"]
    assert abs(res["values"]["this_year"]["mid"] - mid0) < 1e-9 and res["inputs"]["rate"]["this_year"] is None
    print(f"rate: ok (this year's discount rate set 25bp up, the low end at the higher rate: a step of its own "
          f"({move:+.1f} at the mid), the roll-forward's steps and the zero-roll check at last year's; back, the first result)")


def methods_check(eid: int) -> None:
    """This year's value worked out other ways (engine/methods.py), on the same feed: the default is this year's value;
    the recompute ties to the overlay's own formulas; mid-period and mid-year are worth more than end of period, the mid
    at the midpoint rate less than the average of the ends; a method that can't be worked out says why. Preferred, a
    method is this year's value with a bridge step of its own; back to the default, the first result again."""
    import methods
    res = wb.get(eid)["result"]
    inv = res["methods"]
    by = {m["key"]: m for m in inv["methods"]}
    assert [m["key"] for m in inv["methods"]] == [k for k, _l, _w in methods.METHODS] and inv["preferred"] == "overlay", inv
    base, ty = by["overlay"], res["values"]["this_year"]
    assert all(abs(base[e] - ty[e]) < 1e-9 for e in ("low", "mid", "high")), (base, ty)
    assert inv["ties"] and all(abs(by["recompute"]["vs_default"][e]) < 1e-6 for e in ("low", "mid", "high")), by["recompute"]
    assert by["mid_period"]["vs_default"]["mid"] > 0 and by["mid_year"]["vs_default"]["mid"] > 0, (by["mid_period"], by["mid_year"])
    assert by["mid_rate"]["vs_default"]["mid"] < 0 and by["mid_rate"]["low"] is None, by["mid_rate"]
    assert by["overlay_on_date"]["ok"] and by["overlay_on_date"]["vs_default"]["mid"] >= -1e-9, by["overlay_on_date"]
    assert not by["own_flags"]["ok"] and by["own_flags"]["why"], by["own_flags"]  # the pack's overlay has no flags of its own
    mid0, move = ty["mid"], by["mid_period"]["vs_default"]
    finished = lambda: orc.stage(eid, "result").get("finished_at")
    t0 = finished()
    wb.set_method(eid, "mid_period")
    orc.poke()
    wait(eid, lambda v: finished() != t0 and not v["busy"] and status(v)["result"] in orc.SETTLED, "the result by mid-period")
    res = wb.get(eid)["result"]
    assert res["methods"]["preferred"] == "mid_period", res["methods"]["preferred"]
    ch = res["methods"].get("choice") or {}
    assert ch.get("by") == "you" and ch.get("at") and ch.get("previous") is None, ch  # who chose it, when, what it replaced
    try:  # a method this engagement can't work out isn't taken
        wb.set_method(eid, "own_flags")
        raise AssertionError("a method not worked out here was preferred")
    except ValueError as ex:
        assert "can't be worked out here" in str(ex), ex
    for e in ("low", "mid", "high"):
        s = {x["key"]: x for x in res["bridges"][e]["steps"]}
        assert abs(s["method"]["value"] - move[e]) < 1e-6 and abs(res["values"]["this_year"][e] - s["this_year"]["value"]) < 1e-9
    assert abs(res["values"]["this_year"]["mid"] - (mid0 + move["mid"])) < 1e-6
    t0 = finished()
    wb.set_method(eid, None)
    orc.poke()
    wait(eid, lambda v: finished() != t0 and not v["busy"] and status(v)["review"] in orc.SETTLED, "the result back again")
    res = wb.get(eid)["result"]
    assert abs(res["values"]["this_year"]["mid"] - mid0) < 1e-9 and "method" not in [x["key"] for x in res["bridges"]["mid"]["steps"]]
    assert wb.method_choice(eid)["previous"] == "mid_period", wb.method_choice(eid)  # "back to" the one before
    print(f"methods: ok ({sum(m['ok'] for m in inv['methods'])} of {len(inv['methods'])} worked out, the recompute tying to "
          f"the overlay; mid-period {move['mid']:+.1f} at the mid, preferred: a bridge step of its own; back, the first result)")


def review_check(eid: int) -> None:
    """The review is a decision like the others: once per result. A "try again" on the result that works out the same
    result, or on the review itself, doesn't buy another review (the same points stand); a result that changes is
    reviewed again, with the points from before in view."""
    n = lambda: CALLS.count("run_review")
    finished = lambda name: orc.stage(eid, name).get("finished_at")
    before, needs = n(), [x["title"] for x in orc.view(eid)["needs"] if x["stage"] == "review"]
    for name in ("result", "review"):
        t0 = finished(name)
        orc.retry(eid, name)
        wait(eid, lambda v: finished(name) != t0 and not v["busy"] and status(v)["review"] in orc.SETTLED,
             f"the review after trying the {name} again")
        assert n() == before, (name, CALLS[-5:])
        assert [x["title"] for x in orc.view(eid)["needs"] if x["stage"] == "review"] == needs
    t0 = finished("review")
    wb.set_this_year_date(eid, "2026-12-31")  # a different result: reviewed again, its earlier points in view
    orc.poke()
    try:
        wait(eid, lambda v: finished("review") != t0 and status(v)["review"] in orc.SETTLED, "the review of a new result")
        assert n() == before + 1 and "Your points on the run before this one" in PROMPTS["run_review"] and \
            "Terminal-year cash flow" in PROMPTS["run_review"].split("Your points on the run before this one")[1]
    finally:
        t0 = finished("review")
        wb.set_this_year_date(eid, None)
        orc.poke()
        wait(eid, lambda v: finished("review") != t0 and status(v)["review"] in orc.SETTLED, "the review, back again")
    assert n() == before + 1, CALLS[-5:]  # back to the first result: its review stands
    print("review: ok (once per result: a try again on the same result, or on the review, doesn't buy another; a new "
          "result is reviewed again with the earlier points in view)")


NAMES_IN_PACK = ("Asset", "Alpha", "Holdco", "CashFlow", "Val_Inputs", "Operations", "BalanceSheet", "Unlevered",
                 "Recon_PY", "Annual_Summary", "Qtr_Model", "Summary!", "DCF!", "Net cash flow", "valuation_report")


def diagnostics_check(eid: int) -> None:
    """The diagnostics export describes the run in the app's own words, counts, ratios and dates only: nothing of
    the files (no sheet, label or file name), every string known, the models' shapes and how alike they are."""
    import diagnostics
    d = diagnostics.export(eid)
    text = json.dumps(d, default=str).lower()
    leaked = [n for n in NAMES_IN_PACK if n.lower() in text]
    assert not leaked and d["redacted"] == 0, (leaked, d["redacted"])
    m = d["models"]
    assert m["prior_model"]["timelines"]["annual"]["periods_max"] == 21 and "quarterly" not in m["current_model"]["timelines"], m
    assert d["alike"]["prior_model_vs_current_model"]["labels"] > 0.8 and d["profile"]["fy_end_month"] == 6, d["alike"]
    r = d["result"]
    assert r["worked_out"] and r["tied"] == {"low": True, "high": True} and r["terminal"] == "growth_final_year", r
    assert r["discountings"]["low"]["readable"] == 2 and r["gate"]["reliable"] and not r["gate"]["rebuilt"], r
    t = r["gate"]["trace"]  # how the trace did on the rows the value reads, to weigh it on a real model
    assert t["rows"] >= 2 and not t["not_found"] and not t["elsewhere"] and t["confident"] == t["rows"] and \
        t["none"] + t["alone"] + t["agrees"] == t["rows"] and sum(t["scores"].values()) == t["alone"] + t["agrees"], t
    assert (r["gate"]["new_terms"], r["gate"]["new_terms_in_value"], r["gate"]["new_terms_held"],
            r["gate"]["gone_terms"]) == (1, 1, 0, 0), r["gate"]
    print("diagnostics: ok (the run in the app's own words, counts, ratios and dates: no sheet, label or file name; "
          "the models' shapes and how alike they are; how the trace did on the rows the value reads)")


def rebuilt_check() -> None:
    """This year's client model rebuilt from the ground up (its own sheets and labels, quarterly with an annual
    summary, a sheet of last year's figures pasted in under last year's labels, the client's own bridge): the
    value is this year's, not a pasted copy's. The pasted copy is passed over as a candidate for a row and as the
    sheet last year's became (so it doesn't make the horizon look fixed); the rows are found by their numbers; the
    value matches the one this year's ordinary model gives, step for step; a point says the model looks rebuilt and
    lists the rows to check. (The zero-roll check can't catch a pasted copy: it reproduces last year's numbers
    exactly, which is why the copy is refused in the finder.) The diagnostics say all of this without a name."""
    import diagnostics
    ordinary = next(e for e in wb.all_engagements() if e["name"] == "Asset A, FY26")
    want = wb.get(ordinary["id"])["result"]
    e = wb.create("Asset A, FY26 (model rebuilt)")
    eid = e["id"]
    for f in PACK_R:
        upload(eid, f)
    v = wait(eid, lambda v: status(v)["result"] in (*orc.SETTLED, "blocked", "failed") and not v["busy"], "the result", 600)
    res = wb.get(eid)["result"]
    assert abs(res["values"]["this_year"]["mid"] - want["values"]["this_year"]["mid"]) < 1e-6, (res["values"], want["values"])
    assert [(x["key"], round(x["value"], 6)) for x in res["bridges"]["mid"]["steps"]] == \
        [(x["key"], round(x["value"], 6)) for x in want["bridges"]["mid"]["steps"]], res["bridges"]["mid"]["steps"]
    g = res["figures"]["gaps"]
    assert g["rebuilt"] and g["family"] < 0.2 and {x["found"] for x in g["rebuilt_rows"]} == \
        {"Annual_Summary!r6", "Annual_Summary!r7"}, g
    need = next(n for n in v["needs"] if n["id"] == "rebuilt")
    assert need["severity"] == "check" and need["go"]["anchor"] == "rowsCard", need
    info = wb.row_info(eid, "CashFlow!r9")
    assert info["found"] == "Annual_Summary!r6" and info["copies"] == ["Recon_PY!r8"], info
    sess, _ = wb.overlay_session(eid)
    assert sess.rowmap.sheet_for("CashFlow") is None and not sess.fixed_horizon()
    d = diagnostics.export(eid)
    text = json.dumps(d, default=str).lower()
    assert not [n for n in NAMES_IN_PACK if n.lower() in text] and d["redacted"] == 0, d["redacted"]
    cur = d["models"]["current_model"]["timelines"]
    assert cur["quarterly"]["periods_max"] == 80 and cur["annual"]["periods_max"] == 20, cur
    assert d["alike"]["prior_model_vs_current_model"]["labels"] < 0.2 and d["result"]["gate"]["rebuilt"], d["alike"]
    assert {x["kind"] for x in d["result"]["rows"]} == {"rebuilt_rows"} and \
        sum(x["copies_passed_over"] for x in d["result"]["rows"]) == 2, d["result"]["rows"]
    wb.delete(eid)
    print("rebuilt: ok (a rebuilt model with a pasted copy of last year's figures: the copy passed over for rows and "
          "sheets, the rows found by their numbers, this year's value as the ordinary model's, step for step; a point "
          "to check; the diagnostics say so without a name)")


def lines_check(first: int) -> None:
    """Two changes in this year's model the roll can't see, each a point to check that leaves the value as it was: a
    cash-flow line last year's model didn't have, on a sheet the overlay reads (equity injections, with the model's
    own present value of them), and the model saved on another scenario than last year's, with when each was saved
    (this year's by a cell of its own, last year's by the file's properties). The diagnostics count them without a
    name."""
    import diagnostics
    want = wb.get(first)["result"]
    eid = wb.create("Asset A, FY26 (a later model)")["id"]
    confirm_insurance(eid)
    for f in PACK_V:
        upload(eid, f)
    v = wait(eid, lambda v: status(v)["result"] in (*orc.SETTLED, "blocked", "failed") and not v["busy"], "the result", 600)
    res = wb.get(eid)["result"]
    assert status(v)["result"] == "attention" and         abs(res["values"]["this_year"]["mid"] - want["values"]["this_year"]["mid"]) < 1e-6, (status(v), res["values"])
    lines = res["figures"]["gaps"]["new_lines"]
    assert [(x["row"], x["label"], x["periods"], x["total"], x["from"], x["to"], x["units"], (x["pv"] or {}).get("row"))
            for x in lines] == [("CashFlow!r10", "Equity injection", 2, -500.0, "2030-06-30", "2032-06-30", "A$m",
                                 "CashFlow!r11")], lines
    need = next(n for n in v["needs"] if n["id"] == "new-lines")
    assert (need["severity"], need["kind"], need["go"]["anchor"]) == ("check", "check-lines", "linesCard") and         "CashFlow!r10 (Equity injection): 2 period(s), -500.0 in total" in need["detail"], need
    sc = res["scenario"]
    assert [(x["cell"], x["this_year"], x["last_cell"], x["last_year"], x["status"]) for x in sc["selectors"]] ==         [("Inputs!B14", 2.0, "Inputs!B13", 1.0, "differs")], sc["selectors"]
    assert (sc["saved"]["this_year"]["date"], sc["saved"]["this_year"]["from"], sc["saved"]["this_year"]["cell"],
            sc["saved"]["last_year"]["date"], sc["saved"]["last_year"]["from"]) ==         ("2026-07-15", "cell", "Model_Info!B3", "2025-08-14", "file"), sc["saved"]
    need = next(n for n in v["needs"] if n["id"] == "scenario")
    assert (need["severity"], need["kind"], need["go"]["anchor"]) == ("check", "check-scenario", "scenarioCard") and         "it isn't last year's" in need["title"] and "Inputs!B14 (Active scenario (1 base, 2 low, 3 high)) is 2, last "         "year's 1" in need["detail"] and "this year's model was saved 2026-07-15, last year's 2025-08-14" in need["detail"], need
    d = diagnostics.export(eid)
    text = json.dumps(d, default=str).lower()
    names = NAMES_IN_PACK + ("Equity injection", "Model_Info", "Active scenario", "Last saved")
    assert not [n for n in names if n.lower() in text] and d["redacted"] == 0, d["redacted"]
    r = d["result"]
    assert (r["gate"]["new_lines"], r["gate"]["new_line_periods"], r["gate"]["new_lines_error"]) == (1, 2, False), r["gate"]
    assert r["scenario"] == {"found": True, "selectors": 1, "by_status": {"differs": 1}, "error": False,
                             "saved": {"this_year": {"date": "2026-07-15", "from": "cell"},
                                       "last_year": {"date": "2025-08-14", "from": "file"}}}, r["scenario"]
    assert d["run"]["needs"].get("check-lines:check") == 1 and d["run"]["needs"].get("check-scenario:check") == 1, d["run"]
    wb.delete(eid)
    print("lines: ok (a later model's equity injections, on a sheet the overlay reads and in no row of last year's: "
          "2 periods, -500.0, its present value with it; the model saved on scenario 2, last year's on 1, with when "
          "each was saved: two points to check, the value as the ordinary model's; the diagnostics count them)")


def gate_check() -> None:
    """This year's value is held, with a need that says why, where it can't be trusted:
    - a discounting left on last year's valuation date (its date cell not among those the roll moves);
    - the overlay placed as last year's client model with every sheet its own: nothing of this year's model is read
      (placed by a person, the roles stand, and the value waits);
    and the roles check refuses that placement when it's the orchestrator's. A big move at last year's date is a
    point to check, not a hold; a discounting that can't be read makes the roll-forward one step."""
    import overlay as ovmod
    import result
    e = wb.create("Asset A, FY26 (the gate)")
    eid = e["id"]
    confirm_insurance(eid)
    for f in PACK_A:
        upload(eid, f)
    wait(eid, lambda v: status(v)["result"] in orc.SETTLED and status(v)["review"] in orc.SETTLED, "the result")
    # every discounting reads this year's date, or the value waits
    sess, summary = wb.overlay_session(eid)
    roll, cells = summary["roll"], orc.equity_cells(eid)
    assert roll["valuation_date_reads"] and roll["valuation_date_cells"], roll
    g = result.figures(sess, summary, cells)["gaps"]
    assert g["reliable"] and not g["date_cells"]["off"] and not g["no_reads"], g["date_cells"]
    kept = {k: roll[k] for k in ("valuation_date_cells", "valuation_date_cell")}
    roll.update(valuation_date_cells=[], valuation_date_cell=None)
    try:
        g = result.figures(sess, summary, cells)["gaps"]
        assert not g["reliable"] and g["date_cells"]["off"][0]["date"] == "2025-06-30", g["date_cells"]
        st, _, data = orc._result_job(eid, "gate")
        need = next(n for n in data["needs"] if n["id"] == "dates-roll")
        assert st == "blocked" and need["severity"] == "block" and "2025-06-30" in need["detail"], data["needs"]
        assert any(n["id"] == "dates-none" and n["severity"] == "check" for n in data["needs"]), data["needs"]
    finally:
        roll.update(kept)
    # a big move at last year's date: a point to check, the value runs
    band = ovmod.ZERO_ROLL_CHECK
    ovmod.ZERO_ROLL_CHECK = (0.95, 1.05)
    try:
        st, _, data = orc._result_job(eid, "gate")
        need = next(n for n in data["needs"] if n["id"] == "zero-roll")
        assert need["severity"] == "check" and st != "blocked" and "1.10×" in need["title"], data["needs"]
    finally:
        ovmod.ZERO_ROLL_CHECK = band
        orc._result_job(eid, "gate")
    # a discounting that can't be read here: the roll-forward is one step, not its unwind in the new forecast
    t = {"cores": [{"cell": "DCF!C9", "call": "SUMPRODUCT", "inputs": {"rate": 0.08}},
                   {"cell": "DCF!C10", "call": "NPV", "inputs": None}]}
    steps, note = result._steps(None, (t, t["cores"][:1]), 100.0, 130.0, "2026-06-30")
    assert [x["key"] for x in steps] == ["roll"] and steps[0]["value"] == 30.0 and "1 of the 2" in note, (steps, note)
    # the overlay as last year's client model, every sheet its own
    rl = wb.roles(eid)
    ov_id = rl["prior_overlay"]["id"]
    every = next(w["sheet_names"] for w in wb.workbooks(eid) if w["id"] == ov_id)
    a = {k: {"kind": r["kind"], "id": r["id"], "sheets": r["sheets"]} for k, r in rl.items()}
    a.update(prior_model={"kind": "workbook", "id": ov_id, "sheets": None},
             prior_overlay={"kind": "workbook", "id": ov_id, "sheets": every})
    assert "the overlay is inside the client model but takes every sheet: none is left as the client's" in \
        orc.verify_roles(eid, a), orc.verify_roles(eid, a)
    # two roles placed so, the rest left to the orchestrator: it fills in the overlay's sheets, and its check refuses them
    wb.confirm_roles(eid, {"prior_model": a["prior_model"], "prior_overlay": {**a["prior_overlay"], "sheets": None}}, "you")
    orc.poke()
    v = wait(eid, lambda v: status(v)["roles"] == "blocked", "the roles check")
    assert any("takes every sheet" in (h["text"] or "") for h in v["log"] if h["stage"] == "roles"), v["log"][:5]
    # every role placed by a person: checked too, and held, saying why; acknowledged with a reason, their choice
    # stands, and the value waits, saying why
    wb.confirm_roles(eid, {**a, "prior_overlay": {**a["prior_overlay"], "sheets": None}}, "you")
    orc.poke()
    v = wait(eid, lambda v: any(n["id"] == "role-check" for n in v["needs"]), "the check of a person's roles")
    rc = next(n for n in v["needs"] if n["id"] == "role-check")
    assert rc["severity"] == "block" and rc["ack"] and "inside the client model" in rc["detail"], rc
    wb.acknowledge(eid, "role-check", rc["ack"], "set this way on purpose, for the test", rc["title"])
    orc.person(eid, "roles", "acknowledged role-check")
    v = wait(eid, lambda v: any(n["id"] == "no-reads" for n in v["needs"]), "the no-reads hold")
    rc = next(n for n in v["needs"] if n["id"] == "role-check")
    assert rc["severity"] == "info" and rc["acked"]["reason"] == "set this way on purpose, for the test", rc
    need = next(n for n in v["needs"] if n["id"] == "no-reads")
    assert need["severity"] == "block" and need["go"]["anchor"] == "rolesCard", need
    assert wb.get(eid)["result"]["values"]["this_year"] is None, wb.get(eid)["result"]["values"]
    assert all(r["by"] == "you" for k, r in wb.roles(eid).items() if k in ("prior_model", "prior_overlay"))
    wb.delete(eid)
    print("gate: ok (held, with the reason, where a discounting still reads last year's date or nothing of this year's "
          "model is read; the roles check refuses the overlay taking every sheet, a person's roles too until they "
          "acknowledge it with a reason; a big move at last year's date is a "
          "point to check; an unreadable discounting makes the roll-forward one step)")


def workpaper_check(eid: int) -> None:
    import io
    import openpyxl
    import zipfile
    data = workpaper.build(eid)
    charts = [n for n in zipfile.ZipFile(io.BytesIO(data)).namelist() if n.startswith("xl/charts/chart")]
    assert len(charts) == 2, charts  # the bridge's waterfall and the cash flows'
    book = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    want = ["Summary", "Bridge", "Cash flows", "Inputs", "Methods", "Reconciliation", "Key facts", "Files and roles", "Review",
            "Run log"]
    assert book.sheetnames == want, book.sheetnames
    cells = lambda name: [[c for c in row] for row in book[name].iter_rows(values_only=True)]
    res = wb.get(eid)["result"]
    tot = [r for r in cells("Bridge") if r[0] and str(r[0]).startswith("This year, rolled forward")]
    assert tot and abs(tot[-1][2] - res["values"]["this_year"]["mid"]) < 0.05, (tot, res["values"]["this_year"])
    assert not any(re.search(r"\b\d{4}-\d{2}-\d{2}\b", str(c)) for r in cells("Bridge") + cells("Summary") for c in r if c), \
        "a date left as ISO"
    fy = [r for r in cells("Cash flows") if isinstance(r[0], str) and re.fullmatch(r"FY\d{4}", r[0])]
    assert len(fy) == len(res["chart"]["years"]) == 21, len(fy)
    assert any(c and str(c).startswith(workpaper.DISCLAIMER[:40]) for r in cells("Summary") for c in r), "no disclaimer"
    assert not any(re.match(r"gpt-", str(r[3] or "")) for r in cells("Run log")), "a model's name in the log's copy"
    rate = [r for r in cells("Inputs") if r[0] and str(r[0]).startswith("Low (at the higher rate)")]
    assert rate and abs(rate[0][1] - res["inputs"]["rate"]["ends"]["low"]["value"]) < 1e-12 and "!" in rate[0][3], rate
    e = wb.create("Asset A, nothing yet")
    try:
        assert not workpaper.ready(wb.get(e["id"]))
        workpaper.build(e["id"])
        raise AssertionError("a workpaper before the bridge")
    except ValueError:
        pass
    finally:
        wb.delete(e["id"])
    assert workpaper.build(10 ** 6) is None
    print("workpaper: ok (10 sheets and 2 charts; the bridge's mid ends at this year's value; 21 financial years; the "
          "disclaimer; the rate's figure then its cell; roles, not model names; none before the bridge)")


def dates_check() -> None:
    """The image shows a different year from the text layer's valuation date, with the same day: the date is compared
    whole, so it's a difference (its first number, the day, alone would confirm it), and two agreeing reads of the
    image correct it, its YYYYMMDD with it."""
    TRUTH["valuation_date"]["value"] = "30 June 2024"
    try:
        e = wb.create("Asset A, FY26 (the image dates it a year earlier)")
        eid = e["id"]
        for f in PACK_A:
            upload(eid, f)
        wait(eid, lambda v: status(v)["facts"] in (*orc.SETTLED, "failed"), "the facts")
        vd = next(f for f in wb.facts(eid) if f["key"] == "valuation_date")
        assert vd["visual"]["status"] == "corrected" and vd["value_text"] == "30 June 2024" and \
            int(vd["value"]) == 20240630, vd
    finally:
        TRUTH["valuation_date"]["value"] = "30 June 2025"
    print("dates: ok (a date on its image is compared whole, day, month and year; a correction carries its YYYYMMDD)")


def delete_check(first: int) -> None:
    """Deleting an engagement takes everything worked out for it: its reports, rows, overlay folder and call log;
    a model another engagement uses stays until the last of them goes. Then the same files, uploaded again, run from
    the start (every file read afresh) to the same value."""
    import calllog
    want = wb.get(first)["result"]["values"]["this_year"]["mid"]
    eids = [e["id"] for e in wb.all_engagements()]
    files = [library.get(f, full=True) for f in {r["file_id"] for r in wb._q("SELECT file_id FROM eng_files")}]
    docs = [d for e in eids for d in wb._q("SELECT out_dir, source_path FROM documents WHERE engagement_id=?", e)]
    kept_once = False
    for eid in eids:
        t0 = time.time()
        while True:  # the checks before may have left an engagement still running: deleting waits for it
            try:
                r = wb.delete(eid)
                break
            except ValueError as e:
                assert "wait for" in str(e) and time.time() - t0 < 300, e
                time.sleep(0.5)
        kept_once |= bool(r["kept"])
        # every table but stages, which the orchestrator's loop may be writing for it that very moment (an id a
        # new engagement then doesn't take)
        assert wb.get(eid) is None and not (wb.OUT / "overlays" / f"e{eid}").exists() and not any(
            wb._q(f"SELECT 1 FROM {t} WHERE engagement_id=?", eid) for t in ("runlog", "facts", "roles", "eng_files", "documents"))
        with calllog._conn() as db:
            assert not db.execute("SELECT 1 FROM calls WHERE engagement=?", (eid,)).fetchone(), eid
    assert kept_once, "a model several engagements use was deleted with the first of them"
    assert not any(library.get(f["id"]) for f in files) and not any(Path(f["out_dir"]).exists() or
                                                                    Path(f["source_path"]).exists() for f in files)
    assert not any(Path(d["out_dir"]).exists() or Path(d["source_path"]).exists() for d in docs)
    t_gone = time.time()
    again = run_check(PACK_A, "Asset A, FY26 (again, from the start)")
    got = wb.get(again)["result"]["values"]["this_year"]["mid"]
    assert abs(got - want) < 1e-6 and all(w["processed_at"] > t_gone for w in wb.workbooks(again)), (got, want)
    print(f"delete: ok ({len(eids)} engagements deleted with what was worked out for them, a shared model kept until "
          f"the last; the same files run again from the start to {got:,.1f})")


def facts_check() -> None:
    """The code checks behind the models: a figure must be in its quote in the scale and currency it's stated in,
    under its own label, not inside a date or a longer number; a quote on word and number boundaries; a name checked
    as text; a range's two ends; a waiver from the image only for the figure it was given for, and only while the
    quote and figures are the same. The tables: figures under their own row label and column, headings in order.
    The blind image read's context has no figures or dates."""
    import keyfacts as K
    import docingest as D
    import context
    import visual
    fact = lambda **k: {"key": "equity_value", "label": "Equity value", "category": "conclusion", "unit": "A$m",
                        "value_text": "", "low_text": "", "high_text": "", "basis": "", "page": 1, **k}
    failed = lambda f, text: [i["text"] for i in K.check(f, {1: text})["items"] if not i["ok"]]
    # scale and currency
    q = "The equity value of the Target is A$2.3bn on an ex-distribution basis."
    assert failed(fact(value_text="A$2.3m", quote=q), q) == ["A$2.3m: the quote gives it in bn, not m"]
    assert not failed(fact(value_text="A$2.3bn", quote=q), q)
    q = "The equity value is NZ$850m."
    assert failed(fact(value_text="US$850m", quote=q), q) == ["US$850m: the quote gives it in NZ$, not US$"]
    # not the front or the tail of a longer number
    for v, q in (("A$22m", "The equity value is A$22.5m."), ("A$2,296m", "The equity value is A$2,296.7m."),
                 ("2,296.7", "Equity value 2,296.75")):
        assert failed(fact(value_text=v, quote=q), q), (v, q)
    assert K._in_squashed("22", "A$22.0m") and not K._in_squashed("25", "7.25%")
    # under its own label, not another fact's
    q = "WACC of 7.25% and terminal growth of 2.5%"
    rate = fact(key="discount_rate", label="Discount rate", category="assumption", unit="%", value_text="2.5%", quote=q)
    assert failed(rate, q) == ["2.5% sits under 'terminal growth of' in the quote, not under the Discount rate"], failed(rate, q)
    assert not failed({**rate, "key": "terminal_growth_rate", "label": "Terminal growth rate"}, q)
    q = "an equity value of A$2,500m, being an enterprise value of A$4,100m less net debt of A$1,600m"
    assert failed(fact(value_text="A$4,100m", quote=q), q) and not failed(fact(value_text="A$2,500m", quote=q), q)
    q = "Discount rate (post-tax nominal WACC) 7.25% 7.75%"
    assert not failed({**rate, "value_text": "", "low_text": "7.25%", "high_text": "7.75%", "quote": q}, q)
    # franking utilisation as reports put it: a % or a fraction, its label before or after the figure, any rate
    import sourced
    util = lambda v, q: fact(key="franking_utilisation", label="Franking credit utilisation", category="assumption",
                             unit="%", value_text=v, quote=q)
    for v, q in (("80%", "We have ascribed 80% value in franking credits."), ("80%", "an 80% utilisation rate"),
                 ("80%", "a gamma of 0.80"), ("80%", "gamma of 0.8"), ("0.80", "a gamma of 0.80"),
                 ("0.8", "an 80% utilisation rate"), ("60%", "franking credits are utilised at 60%"),
                 ("80%", "WACC of 7.25% with 80% value in franking credits")):
        assert not failed(util(v, q), q), (v, q, failed(util(v, q), q))
    assert failed(util("60%", "an 80% utilisation rate"), "an 80% utilisation rate")
    q = "The value of franking credits is A$389.8m, 15.5% of the equity value."  # its label further back will do
    assert not failed(fact(key="franking_credits_share", label="Franking credits share", unit="%", value_text="15.5%",
                           quote=q), q)
    assert visual._same(util("80%", ""), {"value": "0.80"}) and not visual._same(util("80%", ""), {"value": "0.60"})
    assert sourced._stated([{"key": "franking_utilisation", "value_text": "0.80"}], "franking_utilisation") == [(80.0, "0.80")]
    assert sourced._tie(80.0, (80.0, "0.80")) and sourced._tie(80.0, (80.0, "80%")) and not sourced._tie(60.0, (80.0, "0.8"))
    # the terminal value: how the report works it out, classified in code; what it says, found by searching it
    for text, kind in (("The valuation does not include a terminal value as the concession ends in 2045.", "none"),
                       ("A terminal value is calculated by applying an EV/EBITDA multiple of 12.0x to FY45 EBITDA.", "exit_ebitda"),
                       ("We applied a multiple of 12.0x FY45 EBITDA.", "exit_ebitda"),
                       ("The terminal value reflects an exit multiple of 1.35x the RAB at the end of the forecast.", "exit_rab"),
                       ("We applied an exit multiple of 14.0x.", "exit_other"),
                       ("The terminal cash flow is the average of the FY41 to FY45 cash flows, grown at CPI.", "growth_average"),
                       ("The terminal value applies the Gordon growth model to the final year's normalised cash flow.",
                        "growth_adjusted"),
                       ("A terminal value is calculated using the Gordon growth method.", "growth_final_year"),
                       ("Cash flows are discounted at the cost of equity.", "unknown")):
        got = K.terminal_method([{"key": "terminal_value_method", "value_text": text, "quote": text, "page": 4}])
        assert got["kind"] == kind, (text, got)
    md = ("<!-- page 3 -->\nThe equity value is A$2,507.9m.\n\n<!-- page 4 -->\nA terminal value is calculated using "
          "an exit multiple of 1.35x the RAB. Cash flows are discounted at the cost of equity.\n")
    found = context.terminal(md)
    assert found == [{"page": 4, "text": "A terminal value is calculated using an exit multiple of 1.35x the RAB."}], found
    tv = K.terminal_method([], found)  # no fact for it: the report's own words decide
    assert tv["kind"] == "exit_rab" and tv["from"] == "the report's text" and tv["page"] == 4, tv
    # the fact's words decide over the report's other sentences, and a multiple the valuation implies is a cross-check
    implied = [{"page": 7, "text": "Implied EV/EBITDA multiple: the FY25 multiple of 19.5x to 22.2x implied by our valuation."}]
    said = K.terminal_method([{"key": "terminal_value_method", "value_text": "Gordon growth on FY45 dividends",
                               "quote": "Gordon growth on FY45 dividends", "page": 8}], implied)
    assert said["kind"] == "growth_final_year" and said["from"] == "the fact", said
    assert K.terminal_method([], implied)["kind"] == "unknown"
    assert "not applicable" in sourced.growth_applies(tv) and sourced.growth_applies({"kind": "growth_final_year"}) is None
    f = fact(key="terminal_value_method", label="Terminal value method", category="assumption", unit="text",
             value_text="the average of the FY41 to FY45 cash flows", quote="the average of the FY41 to FY45 cash flows")
    assert not failed(f, "We take the average of the FY41 to FY45 cash flows.") and K.settle_value(dict(f))["value"] is None
    # a date's day isn't a figure; a quote doesn't start inside a number; a name is checked as text
    q = "As at 30 June 2025 the business was valued at A$28m."
    assert failed(fact(value_text="A$30m", quote=q), q)
    assert failed({**rate, "value_text": "25%", "quote": "25%"}, "The WACC of 7.25% applied") == \
        ["quote not found in the document"]
    q = "The Asset 9 Link was valued at A$5m"
    assert failed(fact(key="target_name", label="Target", category="identity", unit="text", value_text="Asset 5 East",
                       quote=q), q)
    # a range: both ends, in one scale; printed as one text, its ends are the low and the high
    q = "The equity value is A$2,100m to A$2,200m"
    assert failed(fact(value_text="A$2,200m", low_text="A$2,100m", quote=q), q) == ["only the low end of the range is given"]
    q = "The equity value is A$2.1bn to A$2,300m"
    assert "the range's ends are in different scales (A$2.1bn and A$2,300m)" in \
        failed(fact(low_text="A$2.1bn", high_text="A$2,300m", quote=q), q)
    q = "The equity value is A$1,900m – A$2,100m (ex-distribution)"
    f = fact(value_text="A$1,900m – A$2,100m", quote=q, basis="ex-distribution", status="approved")
    assert not failed(f, q) and K.settle_value(dict(f))["value"] is None
    c = K.conclusion([f])
    assert (c["low"], c["mid"], c["high"]) == (1900.0, 2000.0, 2100.0) and c["texts"]["high_text"] == "A$2,100m", c
    assert K.numbers(K._unrange("Net debt at valuation - 850.0")) == ["-850.0"]
    # the image's waiver: a figure's check only, for that quote and those figures
    q = "The equity value is A$2,507.9m."
    f = fact(value_text="A$2,507.9m", quote=q)
    visual._apply(f, {"value": "A$2,570.9m"}, {1: q}, "two reads of the image agree")
    assert f["check"]["ok"] and "waived" in " ".join(i["text"] for i in f["check"]["items"]), f["check"]
    moved = {1: q + " Its equity value is A$2,507.9m, as before."}  # another quote on the page: checked afresh
    assert failed({**f, "quote": "Its equity value is A$2,507.9m, as before."}, moved[1]) == \
        ["A$2,570.9m is not in the quote"]
    f = fact(value_text="A$2,507.9m", quote="words on no page")
    visual._apply(f, {"value": "A$2,507.9m"}, {1: q}, "two reads of the image agree")
    assert not f["check"]["ok"] and "quote not found in the document" in [i["text"] for i in f["check"]["items"]
                                                                            if not i["ok"]], f["check"]
    # tables: figures under their own label and column, headings in order
    lines = ["A$m Low High", "Equity value (ex-distribution) 2,296.7 2,500.0", "Equity value (cum-distribution) 2,350.1 2,553.4"]
    t = lambda h, ex, cum: f"| A$m | {h} |\n|---|---|---|\n| Equity value (ex-distribution) | {ex} |\n" \
                           f"| Equity value (cum-distribution) | {cum} |"
    assert D.check_text_layer(t("Low | High", "2,296.7 | 2,500.0", "2,350.1 | 2,553.4"), lines)["ok"]
    assert not D.check_text_layer(t("Low | High", "2,350.1 | 2,553.4", "2,296.7 | 2,500.0"), lines)["ok"]  # rows swapped
    assert not D.check_text_layer(t("High | Low", "2,296.7 | 2,500.0", "2,350.1 | 2,553.4"), lines)["ok"]  # headings
    assert D.check_text_layer("| A$m | Low | High |\n|---|---|---|\n| Equity value (ex-distribution) | 2,296.7 | 2,500.0 |",
                              ["A$m Low High", "Equity value", "(ex-distribution) 2,296.7 2,500.0"])["ok"]  # two lines
    one = "| A$m | Low | High |\n|---|---|---|\n| Equity value | 2,296.7 | 2,500.0 |"
    assert D.compare_reads(one, one)["ok"]
    for other in ("| A$m | Low | High |\n|---|---|---|\n| Equity value | | 2,296.7 | 2,500.0 |",
                  "| A$m | Low | High |\n|---|---|---|\n| Equity value | 2,500.0 | 2,296.7 |",
                  "| A$m | High | Low |\n|---|---|---|\n| Equity value | 2,296.7 | 2,500.0 |"):
        assert not D.compare_reads(one, other)["ok"], other
    # the blind read's context: what things are, not the figures or the date
    md = ("<!-- page 1 -->\nDear Board, we have assessed the equity value of Asset A Pty Ltd as at 30 June 2025 at "
          "A$2,507.9m (ex-distribution).\n\n<!-- page 2 -->\nScope: the equity value of Asset A.\n")
    facts = [{"key": "valuation_date", "value_text": "30 June 2025"}, {"key": "currency_units", "value_text": "A$m"}]
    full, blind = (context.block(md, page=1, facts=facts, blind=b) for b in (False, True))
    assert "2,507.9" in full and "30 June 2025" in full, full
    assert "2,507.9" not in blind and "June" not in blind and "A$m" in blind and "Asset A" in blind, blind
    print("facts: ok (a figure in its quote's scale and currency, under its own label, not in a date or a longer number; "
          "names as text; a range's ends; the image waives only a figure, for its quote; table rows and headings in "
          "place; the blind read sees no figures)")


def ranges_check() -> None:
    import keyfacts
    import context
    import docingest
    for kind in ("pdf", "pptx"):
        d = docingest.process(str(PACK / f"AssetA_valuation_report_FY25.{kind}"), tempfile.mkdtemp(), read=False)
        c = context.find(d["markdown"])
        assert c["letter"]["pages"] == [2] and (kind == "pptx" or c["scope"]["pages"] == [2]), (kind, c)
    n = lambda t: keyfacts.numbers(keyfacts._unrange(t))
    assert n("Discount rate 7.25% - 7.75% Post-tax") == ["7.25%", "7.75%"]
    assert n("7.25%–7.75%") == ["7.25%", "7.75%"]
    assert n("Less: net debt (850.0) (850.0)") == ["-850.0", "-850.0"]
    assert n("Net debt at valuation - 850.0") == ["-850.0"], n("Net debt at valuation - 850.0")
    got = n("FY25 - 30.0")
    assert got == ["25", "30.0"], got  # a year then a figure reads as a range: the check only looks for the fact's own numbers
    import visual
    from datetime import date
    for t in ("as at 30 June 2025 (the", "30 Jun 25", "30-Jun-25", "June 30, 2025", "30/06/2025", "2025-06-30",
              "30th of June 2025"):
        assert keyfacts.date_of(t) == date(2025, 6, 30), t
    assert keyfacts.date_of("30 and 2025") is None and keyfacts.date_of("FY25 7.25%") is None
    vd = {"key": "valuation_date", "unit": "date", "value_text": "30 June 2025"}
    assert visual._same(vd, {"value": "30/06/2025"}) and not visual._same(vd, {"value": "30 September 2025"}) and \
        not visual._same(vd, {"value": "30 June 2024"}) and not visual._same(vd, {"value": ""})
    got = keyfacts.settle_value({**vd, "value": 20240630})["value"]
    assert got == 20250630, got  # the model's YYYYMMDD gives way to the date in the text
    chk = keyfacts.check({**vd, "value_text": "30 July 2025", "value": 20250730, "page": 2,
                          "quote": "as at 30 June 2025 (the"}, {2: "as at 30 June 2025 (the"})
    assert not chk["ok"] and "30 July 2025 is not a date in the quote" in [i["text"] for i in chk["items"]], chk
    print("ranges: ok (a dash between figures is a range; a bracket, or a dash after a word, is a minus; the letter and "
          "the scope are found; a date is read and checked whole)")


def changes_check() -> None:
    """The two finders behind the lines and the scenario, on two small models: the cash-flow lines the overlay
    doesn't read are this year's new equity injections (their present value with them) and a shareholder loan line
    last year's model had empty, not the contributions the row the overlay reads adds up, the distributions last
    year's model had too, or a cost of equity; the selectors are a number, an option's name and a switch the
    formulas read, not a note nothing reads, a series, or a case named in passing ("Base case capex"); this year's
    save time is the model's own cell, last year's the file's."""
    from datetime import date, datetime
    from types import SimpleNamespace

    import xlsxwriter
    from xlsxwriter.utility import xl_col_to_name as col
    import build_map
    import overlay as ovmod
    import result
    import rowfind
    import scenarios
    out = Path(tempfile.mkdtemp(prefix="changes_"))

    def book(name, fy0, rows, scen, saved=None):
        path = out / f"{name}.xlsx"
        w = xlsxwriter.Workbook(path)
        w.set_properties({"created": datetime(2024, 12, 3)})
        dt = w.add_format({"num_format": "dd-mmm-yy"})
        cf = w.add_worksheet("CF")
        cf.write(2, 0, "Period ending")
        ends = [date(fy0 + k, 6, 30) for k in range(10)]
        for k, e in enumerate(ends):
            cf.write_datetime(2, 2 + k, e, dt)
        for r, (label, vals) in rows.items():
            cf.write(r - 1, 0, label)
            for k, e in enumerate(ends):
                v = vals(e, k)  # a number, or a formula and the value Excel saved for it
                if isinstance(v, tuple):
                    cf.write_formula(r - 1, 2 + k, v[0].replace("#", col(2 + k)), None, v[1])
                else:
                    cf.write_number(r - 1, 2 + k, v)
        sc = w.add_worksheet("Scen")
        sc.write(2, 0, "Select scenario"); sc.write(2, 1, scen[0])
        sc.write(3, 0, "Scenario in use"); sc.write_formula(3, 1, '=CHOOSE(B3,"Base","Low","High")')
        sc.write(5, 0, "Case"); sc.write(5, 2, scen[1])
        sc.write(6, 0, "Case flag"); sc.write_formula(6, 1, '=IF(C6="Downside",1,0)')
        sc.write(8, 0, "Sensitivity switch: capex"); sc.write(8, 1, 0)
        sc.write(9, 0, "Capex with sensitivity"); sc.write_formula(9, 1, "=100*(1+0.1*B9)")
        sc.write(11, 0, "Scenario notes"); sc.write(11, 1, 1)
        sc.write(13, 0, "Base case capex"); sc.write(13, 1, 15)
        sc.write(14, 0, "Capex"); sc.write_formula(14, 1, "=B14*2")
        sc.write(16, 0, "Scenario volumes")
        for k in range(5):
            sc.write(16, 1 + k, k + 1)
        sc.write(17, 0, "Total volume"); sc.write_formula(17, 1, "=SUM(B17:F17)")
        if saved:
            info = w.add_worksheet("Info")
            info.write(2, 0, "Last saved"); info.write_datetime(2, 1, saved, w.add_format({"num_format": "dd-mmm-yy hh:mm"}))
        w.close()
        return path, ovmod.Workbook(build_map.main(str(path), str(out / name))["db"])

    common = {5: ("Revenue", lambda e, k: 100.0 + k), 6: ("Equity contributions", lambda e, k: -10.0),
              7: ("Distributions paid", lambda e, k: 30.0 + k), 8: ("Net cash flow to equity", lambda e, k: ("=#5+#6", 90.0 + k)),
              10: ("Cost of equity", lambda e, k: 0.11)}
    p_path, prior = book("prior", 2026, {**common, 9: ("Shareholder loan drawdowns", lambda e, k: 0.0)}, (3, "Base"))
    injection = {2030: -225.0, 2032: -275.0}
    c_path, cur = book("current", 2027, {**common, 9: ("Shareholder loan drawdowns", lambda e, k: -20.0),
                                         11: ("Equity injection", lambda e, k: injection.get(e.year, 0.0)),
                                         12: ("PV of equity injection", lambda e, k: (f"=#11/1.1^{k + 1}",
                                                                                     injection.get(e.year, 0.0) / 1.1 ** (k + 1)))},
                       (2, "Downside"), saved=datetime(2025, 11, 26, 14, 5))
    sess = SimpleNamespace(current=cur, prior=prior, ov=prior, sheets=[], client_sheets=set(), client_link=None,
                           ext_cached={}, rowmap=rowfind.RowFinder(ovmod.RowMap(prior, cur), prior, cur))
    lines = result._new_lines(sess, [("CF", 8)], "2026-06-30", "2025-06-30")
    assert [(x["row"], x["periods"], round(x["total"], 6), (x["pv"] or {}).get("row"), x["why"].split(" ")[0])
            for x in lines] == [("CF!r11", 2, -500.0, "CF!r12", "no"), ("CF!r9", 10, -200.0, None, "last")], lines
    got = scenarios.settings(sess, {"wiring": {"current": {"source_path": str(c_path)}, "prior": {"source_path": str(p_path)}}})
    assert [(x["cell"], x["this_year"], x["last_year"], x["status"]) for x in got["selectors"]] == \
        [("Scen!B3", 2.0, 3.0, "differs"), ("Scen!C6", "Downside", "Base", "differs"), ("Scen!B9", 0.0, 0.0, "same")], got
    assert (got["saved"]["this_year"]["date"], got["saved"]["this_year"]["cell"], got["saved"]["last_year"]["date"],
            got["saved"]["last_year"]["from"]) == ("2025-11-26", "Info!B3", "2024-12-03", "file"), got["saved"]
    for w in (prior, cur):
        w.close()
    print("changes: ok (the lines the overlay doesn't read: a new equity injection with its present value, a loan line "
          "empty last year; not what the read row adds up, last year's distributions or a rate; the selectors a number, "
          "an option and a switch the formulas read, not a note, a series or a case in passing; when each was saved)")


def terms_check() -> None:
    """The row finder's trace and the new terms, on two small models. This year's tax is on a new sheet under a label
    with no word of last year's, and at another rate (its numbers aren't last year's): the trace up from EBITDA and
    capital expenditure (anchors: the same label, last year's numbers) reaches it as the one row left once the
    others are paired, so it's found. This year's cash flow available adds lease payments and drops last year's working
    capital movement. Read through the distributions the overlay reads, both move the value and hold it until
    confirmed (the new one belongs, the old one's gone); with the overlay reading the cash flow's terms instead,
    neither does, each a point to check."""
    from datetime import date
    from types import SimpleNamespace

    import xlsxwriter
    from xlsxwriter.utility import xl_col_to_name as col
    import build_map
    import overlay as ovmod
    import result
    import rowfind
    out = Path(tempfile.mkdtemp(prefix="terms_"))

    def book(name, fy0, rate, grow, moved):
        path = out / f"{name}.xlsx"
        w = xlsxwriter.Workbook(path)
        dt = w.add_format({"num_format": "dd-mmm-yy"})
        inp = w.add_worksheet("Inputs")
        inp.write(2, 0, "Tax rate")
        inp.write(2, 1, rate)
        sheets = [w.add_worksheet("CF")] + ([w.add_worksheet("Taxation")] if moved else [])
        ends = [date(fy0 + k, 6, 30) for k in range(10)]
        for sh in sheets:
            sh.write(2, 0, "Period ending")
            for k, e in enumerate(ends):
                sh.write_datetime(2, 2 + k, e, dt)
        cf = sheets[0]
        rows = [(5, "Revenue"), (6, "Operating costs"), (7, "EBITDA"), (8, "Capital expenditure"),
                (9, "Lease payments" if moved else "Tax paid"), (10, "Cash flow available"), (12, "Distributions to equity")]
        if not moved:
            rows.append((11, "Working capital movement"))
        for r, label in rows:
            cf.write(r - 1, 0, label)
        if moved:
            sheets[1].write(5, 0, "Cash taxes")
        for k in range(10):
            c = col(2 + k)
            rev, opex, capex = (200.0 + 5 * k) * grow, -(80.0 + 2 * k) * grow, -20.0 * grow
            tax = -(rev + opex + capex) * rate
            lease = -12.0 if moved else 0.0
            cf.write_number(f"{c}5", rev)
            cf.write_number(f"{c}6", opex)
            cf.write_formula(f"{c}7", f"={c}5+{c}6", None, rev + opex)
            cf.write_number(f"{c}8", capex)
            if moved:
                sheets[1].write_formula(f"{c}6", f"=-(CF!{c}7+CF!{c}8)*Inputs!$B$3", None, tax)
                cf.write_number(f"{c}9", lease)
                cf.write_formula(f"{c}10", f"={c}7+{c}8+Taxation!{c}6+{c}9", None, rev + opex + capex + tax + lease)
            else:
                cf.write_formula(f"{c}9", f"=-({c}7+{c}8)*Inputs!$B$3", None, tax)
                cf.write_number(f"{c}11", -5.0)
                cf.write_formula(f"{c}10", f"={c}7+{c}8+{c}9+{c}11", None, rev + opex + capex + tax - 5.0)
            wc = 0.0 if moved else -5.0
            cf.write_formula(f"{c}12", f"={c}10*0.9", None, 0.9 * (rev + opex + capex + tax + lease + wc))
        w.close()
        return ovmod.Workbook(build_map.main(str(path), str(out / name))["db"])

    prior, cur = book("prior", 2026, 0.30, 1.0, False), book("current", 2027, 0.40, 1.03, True)
    rm = rowfind.RowFinder(ovmod.RowMap(prior, cur), prior, cur)
    ex = rm.explain("CF", 9)
    assert ex["found"] == ("Taxation", 6) and "trace" in dict(ex["evidence"]) and rm.confident("CF", 9), ex
    plain = rowfind.RowFinder(ovmod.RowMap(prior, cur), prior, cur)
    plain._by_trace = lambda s, r: []
    assert not plain.confident("CF", 9), plain.explain("CF", 9)  # without the trace it isn't settled
    assert rm.terms("CF", 10) == ([("CF", 9)], [("CF", 11)]) and rm.terms("CF", 9) == ([], []) and \
        rm.terms("CF", 12) == ([], []), (rm.terms("CF", 10), rm.terms("CF", 9), rm.terms("CF", 12))
    sess = SimpleNamespace(current=cur, prior=prior, ov=prior, rowmap=rm)
    vds = ("2026-06-30", "2025-06-30")
    new, gone = result._term_changes(sess, [("CF", 12)], *vds, set(), set())
    assert [(x["row"], x["under"], x["in_value"], x["hold"], x["periods"], round(x["total"], 6)) for x in new] == \
        [("CF!r9", "CF!r10", True, True, 10, -120.0)], new
    assert [(x["key"], x["last_year"], x["in_value"], x["hold"], x["periods"], x["total"], x["now"]) for x in gone] == \
        [("was:CF!r11", "CF!r10", True, True, 10, -50.0, None)], gone
    new, gone = result._term_changes(sess, [("CF", 12)], *vds, {"CF!r9", "was:CF!r11"}, set())
    assert not new[0]["hold"] and not gone[0]["hold"] and gone[0]["confirmed"], (new, gone)
    new, gone = result._term_changes(sess, [("CF", 7), ("CF", 8), ("CF", 9)], *vds, set(), set())
    assert [(x["key"], x["in_value"], x["hold"]) for x in new + gone] == \
        [("CF!r9", False, False), ("was:CF!r11", False, False)], (new, gone)
    assert not result._term_changes(sess, [("CF", 7), ("CF", 8), ("CF", 9)], *vds, set(), {"CF!r9"})[0]
    for w in (prior, cur):
        w.close()
    print("terms: ok (tax moved to a new sheet under another label and at another rate, found by the trace up from "
          "EBITDA and capex, and not without it; lease payments new in the cash flow available, its working capital "
          "gone: through the distributions both move the value, holding it until confirmed; beside the terms the "
          "overlay reads, flagged)")


def upgrade_check() -> None:
    """A database from an earlier version gets the columns added since (a Windows install upgrades in place)."""
    import sqlite3
    db_path = Path(tempfile.mkdtemp()) / "old.db"
    with sqlite3.connect(db_path) as old:
        old.execute("CREATE TABLE facts(id INTEGER PRIMARY KEY, engagement_id INT, key TEXT)")
    keep = wb.DB
    wb.DB = db_path
    try:
        cols = {r[1] for r in wb._conn().execute("PRAGMA table_info(facts)")}
    finally:
        wb.DB = keep
    assert {"visual_json", "agent_json", "status"} <= cols, cols
    print("upgrade: ok (an older database gets the columns added since)")


def main() -> None:
    ranges_check()
    facts_check()
    changes_check()
    terms_check()
    upgrade_check()
    if not PACK.exists():
        sys.exit("run tests/make_pack.py first")
    sandbox()
    stub_models()
    library.start()
    orc.start()
    eid = run_check()
    workpaper_check(eid)
    overview_check(eid)
    diagnostics_check(eid)
    rebuilt_check()
    lines_check(eid)
    gate_check()
    review_check(eid)
    held_check(eid)
    rate_check(eid)
    methods_check(eid)
    rows_context_check(eid)
    gating_check(eid)
    roles_check()
    adviser_tabs_check()
    copies_check()
    escalate_check()
    dates_check()
    place_check()
    other = run_check(PACK_B, "Asset A, FY26 (overlay inside)")
    a, b = (wb.get(x)["result"]["values"]["this_year"]["mid"] for x in (eid, other))
    assert abs(a - b) < 1e-6, f"the same files give different values in the two layouts: {a} vs {b}"
    print("layouts: ok (the overlay standalone and inside the client model give the same value this year: the roll "
          "moves the valuation date the discountings read)")
    delete_check(eid)


if __name__ == "__main__":
    main()
