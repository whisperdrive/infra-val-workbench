"""Write a synthetic recurring-valuation pack to tests/pack/ (git-ignored; re-run to regenerate).

A fictional infrastructure asset ("Asset A", code name Project Alpha), valued a year ago and now being rolled
forward. Two layouts of the same prior valuation, so both ways an overlay is delivered are covered:

  Pack A (overlay standalone)                  Pack B (overlay inside the client model)
    AssetA_valuation_report_FY25.pdf             AssetA_valuation_report_FY25.pptx
    AssetA_BP25_client_model_Jun25.xlsx          AssetA_BP25_with_overlay.xlsx
    Alpha_valuation_overlay_FY25.xlsx            (Val_Inputs / DCF / Summary sheets inside it)
    AssetA_BP26_client_model_Jun26.xlsx          AssetA_BP26_client_model_Jun26.xlsx
    AssetA_FY26_plan_rebuilt.xlsx (this year's model rebuilt: other sheets and labels, quarterly, a lookalike
                                   sheet of last year's numbers, the client's own bridge)

The valuation follows the conventions the workbench is built for:
  - the conclusion is the equity value as a low / mid / high range: low at the higher discount rate, high at
    the lower one, and the mid is the midpoint of the two (not the value at the middle rate)
  - ex-distribution: the declared distribution comes off; the cum-distribution value is quoted once, in passing
  - the value of franking credits (tax paid x utilisation, discounted like the cash flows) is part of it
  - the report discloses the terminal value and the present values of the forecast and of the terminal value
In the PDF the key assumptions table (the discount rate is only there) is pasted as a picture, so it has to be read
from the image; in the slides every table is the slide's own. The standalone overlay reads the client model through
real external links ([1]CashFlow!D9, with the link's
cached values), as Excel saves them. Every number the report quotes is computed here from the models, so
report, overlay and client model agree.
    uv run python tests/make_pack.py
"""
import io
import re
import sys
import zipfile
from datetime import date
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import xlsxwriter  # noqa: E402
from matplotlib.backends.backend_pdf import PdfPages  # noqa: E402
from xlsxwriter.utility import xl_col_to_name as COL  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "engine"))
import dcf  # noqa: E402

OUT = ROOT / "tests" / "pack"
YEARS = 20
FIRST_COL = 3  # column D
RATES = (0.0775, 0.0725)  # the low value's rate (higher), the high value's (lower)
GROWTH, NET_DEBT, DISTRIBUTION, UTILISATION = 0.025, 850.0, 25.0, 0.50
# this year's client model's balances at its own valuation date: what the overlay's typed net debt and declared
# distribution (last year's, held in the roll-forward) would be this year
NET_DEBT_NEW, DISTRIBUTION_NEW = 880.0, 27.5
VD, VD_NEW = date(2025, 6, 30), date(2026, 6, 30)


# ---- the client model ------------------------------------------------------------------------------------

def client_numbers(fy0: int, volume0: float, growth: float, tariff0: float, cpi: float, opex_pct: float,
                   capex: float, major: float, insurance: float | None) -> dict:
    ends = [date(fy0 + i, 6, 30) for i in range(YEARS)]
    volume = [volume0 * (1 + growth) ** i for i in range(YEARS)]
    tariff = [tariff0 * (1 + cpi) ** i for i in range(YEARS)]
    revenue = [v * p for v, p in zip(volume, tariff)]
    opex = [-r * opex_pct for r in revenue]
    ins = [-(insurance or 0) * (1 + cpi) ** i for i in range(YEARS)]
    ebitda = [r + o + s for r, o, s in zip(revenue, opex, ins)]
    cap = [-(capex * (1 + cpi) ** i + (major if (fy0 + i) % 5 == 0 else 0)) for i in range(YEARS)]
    tax = [-max(0.0, (e + c) * 0.30) for e, c in zip(ebitda, cap)]
    fcf = [e + c + t for e, c, t in zip(ebitda, cap, tax)]
    return dict(ends=ends, volume=volume, tariff=tariff, revenue=revenue, opex=opex, insurance=ins, ebitda=ebitda,
                capex=cap, tax=tax, fcf=fcf)


def write_client(wb, n: dict, inputs: dict, insurance: bool, balances: tuple[float, float] | None = None) -> dict:
    """Inputs / Operations / CashFlow sheets, and a BalanceSheet whose first column is the opening balance at the
    model's valuation date (net debt and the distribution payable: balances). Returns {line item: (sheet, row)}
    (1-based rows)."""
    b = wb.add_format({"bold": True})
    pct, num, dt = wb.add_format({"num_format": "0.00%"}), wb.add_format({"num_format": "#,##0.0"}), \
        wb.add_format({"num_format": "dd-mmm-yy"})
    i = wb.add_worksheet("Inputs")
    i.write(0, 0, "Asset A - business plan inputs", b)
    at = {}
    for r, (label, value, unit) in enumerate(inputs["rows"], start=2):
        i.write(r, 0, label)
        if isinstance(value, date):
            i.write_datetime(r, 1, value, dt)
        else:
            i.write(r, 1, value, pct if unit == "%" else num)
        i.write(r, 2, unit)
        at[label] = f"Inputs!$B${r + 1}"

    def timeline(ws, title):
        ws.write(0, 0, title, b)
        ws.write(2, 0, "Period ending", b)
        for k, d in enumerate(n["ends"]):
            ws.write_datetime(2, FIRST_COL + k, d, dt)
        ws.write(3, 0, "Financial year")
        for k, d in enumerate(n["ends"]):
            ws.write(3, FIRST_COL + k, f"FY{d.year % 100:02d}")

    rows = {}
    o = wb.add_worksheet("Operations")
    timeline(o, "Operations")
    op_rows = [("Volume", "m units", "volume"), ("Tariff (nominal)", "A$", "tariff"), ("Usage income", "A$m", "revenue"),
               ("Operating costs", "A$m", "opex")]
    if insurance:
        op_rows.append(("Insurance", "A$m", "insurance"))
    op_rows.append(("EBITDA", "A$m", "ebitda"))
    for k, (label, unit, key) in enumerate(op_rows):
        rows[key] = ("Operations", 6 + k)
        o.write(5 + k, 0, label)
        o.write(5 + k, 1, unit)
    rr = {key: r for key, (_, r) in rows.items()}
    for k in range(YEARS):
        c, p = COL(FIRST_COL + k), COL(FIRST_COL + k - 1)
        o.write_formula(f"{c}{rr['volume']}", f"={at['Opening volume']}" if k == 0 else
                        f"={p}{rr['volume']}*(1+{at['Volume growth']})", num, n["volume"][k])
        o.write_formula(f"{c}{rr['tariff']}", f"={at['Tariff at start of forecast']}" if k == 0 else
                        f"={p}{rr['tariff']}*(1+{at['CPI']})", num, n["tariff"][k])
        o.write_formula(f"{c}{rr['revenue']}", f"={c}{rr['volume']}*{c}{rr['tariff']}", num, n["revenue"][k])
        o.write_formula(f"{c}{rr['opex']}", f"=-{c}{rr['revenue']}*{at['Operating costs (% of revenue)']}", num,
                        n["opex"][k])
        parts = f"{c}{rr['revenue']}+{c}{rr['opex']}"
        if insurance:
            o.write_formula(f"{c}{rr['insurance']}", f"=-{at['Insurance premium']}*(1+{at['CPI']})^{k}", num,
                            n["insurance"][k])
            parts += f"+{c}{rr['insurance']}"
        o.write_formula(f"{c}{rr['ebitda']}", f"={parts}", num, n["ebitda"][k])

    cf = wb.add_worksheet("CashFlow")
    timeline(cf, "Cash flow")
    cf_rows = [("EBITDA", "ebitda"), ("Capital expenditure", "capex"), ("Tax paid", "tax"),
               ("Unlevered free cash flow", "fcf")]
    for k, (label, key) in enumerate(cf_rows):
        cf.write(5 + k, 0, label)
        cf.write(5 + k, 1, "A$m")
        rows[f"cf_{key}"] = ("CashFlow", 6 + k)
    e_row = rows["ebitda"][1]
    for k in range(YEARS):
        c = COL(FIRST_COL + k)
        fy = n["ends"][k].year
        cf.write_formula(f"{c}6", f"=Operations!{c}{e_row}", num, n["ebitda"][k])
        major = f"+IF(MOD({fy},5)=0,{at['Major maintenance (every 5 years)']},0)"
        cf.write_formula(f"{c}7", f"=-({at['Maintenance capex']}*(1+{at['CPI']})^{k}{major})", num, n["capex"][k])
        cf.write_formula(f"{c}8", f"=-MAX(0,({c}6+{c}7)*{at['Tax rate']})", num, n["tax"][k])
        cf.write_formula(f"{c}9", f"={c}6+{c}7+{c}8", num, n["fcf"][k])
    if balances:
        bs = wb.add_worksheet("BalanceSheet")
        bs.write(0, 0, "Balance sheet (extract)", b)
        bs.write(2, 0, "Balance at", b)
        opening = date(n["ends"][0].year - 1, 6, 30)
        for k, d in enumerate([opening] + n["ends"]):
            bs.write_datetime(2, FIRST_COL - 1 + k, d, dt)
        for r, (label, v0, step) in enumerate((("Net debt", balances[0], 0.97), ("Distribution payable", balances[1], 1.03)), 5):
            bs.write(r, 0, label)
            bs.write(r, 1, "A$m")
            for k in range(YEARS + 1):
                bs.write_number(r, FIRST_COL - 1 + k, round(v0 * step ** k, 1), num)
            rows[label] = ("BalanceSheet", r + 1)
    return rows


# ---- this year's client model, rebuilt ------------------------------------------------------------------------

def write_rebuilt(wb, n: dict, prior: dict, vd: date) -> None:
    """This year's client model rebuilt from the ground up, with this year's economics (n) and almost nothing of
    last year's to go by: its own sheet names and labels, quarterly cash flows (each year's a quarter at a time) with
    an annual summary that totals them, a sheet of last year's numbers pasted in under last year's labels (a
    lookalike: last year's figures, not this year's), the client's own bridge from last year's value, and its
    balances under other names. The value it should give is the one this year's ordinary model gives."""
    b = wb.add_format({"bold": True})
    num, dt, pct = wb.add_format({"num_format": "#,##0.0"}), wb.add_format({"num_format": "dd-mmm-yy"}), \
        wb.add_format({"num_format": "0.00%"})
    a = wb.add_worksheet("Assumptions")
    a.write(0, 0, "Key assumptions", b)
    for r, (label, v, f) in enumerate((("Valuation date", vd, dt), ("Annual escalation", 0.03, pct),
                                       ("Corporate tax rate", 0.30, pct)), start=2):
        a.write(r, 0, label)
        (a.write_datetime if isinstance(v, date) else a.write)(r, 1, v, f)
    quarters = [date(e.year - 1, 9, 30) if q == 0 else date(e.year - 1, 12, 31) if q == 1 else date(e.year, 3, 31) if q == 2
                else e for e in n["ends"] for q in range(4)]
    qm = wb.add_worksheet("Qtr_Model")
    qm.write(0, 0, "Quarterly cash flow model", b)
    qm.write(2, 0, "Quarter ending", b)
    for k, d in enumerate(quarters):
        qm.write_datetime(2, 2 + k, d, dt)
    parts = (("Gross receipts", "revenue", 1), ("Operating spend", "opex", 1), ("Insurance cost", "insurance", 1),
             ("Sustaining capex", "capex", 1), ("Income tax", "tax", 1))
    for i, (label, key, sign) in enumerate(parts):
        qm.write(5 + i, 0, label)
        for k in range(len(quarters)):
            qm.write_number(5 + i, 2 + k, sign * n[key][k // 4] / 4, num)
    net = 5 + len(parts)
    qm.write(net, 0, "Net cash flow", b)
    for k in range(len(quarters)):
        c = COL(2 + k)
        qm.write_formula(net, 2 + k, f"=SUM({c}6:{c}{net})", num,
                         sum(n[key][k // 4] / 4 for _l, key, _s in parts))
    an = wb.add_worksheet("Annual_Summary")
    an.write(0, 0, "Annual summary", b)
    an.write(2, 0, "Year ending", b)
    for k, e in enumerate(n["ends"]):
        an.write_datetime(2, FIRST_COL + k, e, dt)
    for i, (label, row, key) in enumerate((("Net cash flow (FY)", net, "fcf"), ("Income tax (FY)", 9, "tax"))):
        an.write(5 + i, 0, label)
        for k in range(YEARS):
            q0, q3 = COL(2 + 4 * k), COL(2 + 4 * k + 3)
            an.write_formula(5 + i, FIRST_COL + k, f"=SUM(Qtr_Model!{q0}{row + 1}:{q3}{row + 1})", num, n[key][k])
    rc = wb.add_worksheet("Recon_PY")  # last year's numbers, pasted in under last year's labels
    rc.write(0, 0, "Prior-year model values (pasted for reconciliation)", b)
    rc.write(2, 0, "Period ending", b)
    for k, e in enumerate(prior["ends"]):
        rc.write_datetime(2, FIRST_COL + k, e, dt)
    for i, (label, key) in enumerate((("EBITDA", "ebitda"), ("Tax paid", "tax"), ("Unlevered free cash flow", "fcf"))):
        rc.write(5 + i, 0, label)
        for k in range(YEARS):
            rc.write_number(5 + i, FIRST_COL + k, prior[key][k], num)
    vb = wb.add_worksheet("Val_Bridge")  # the client's own bridge from last year's value
    vb.write(0, 0, "Valuation movement (client)", b)
    for i, (label, v) in enumerate((("Prior valuation", 2507.9), ("Unwind of discount", 253.3), ("Cash flows received", -160.5),
                                    ("Forecast changes", 784.0), ("Current valuation", 3384.7))):
        vb.write(2 + i, 0, label)
        vb.write_number(2 + i, 1, v, num)
    nb = wb.add_worksheet("Net_Borrowings")
    nb.write(0, 0, "Funding position", b)
    nb.write(2, 0, "As at", b)
    for k, d in enumerate([vd] + n["ends"][:5]):
        nb.write_datetime(2, FIRST_COL - 1 + k, d, dt)
    for i, (label, v0) in enumerate((("Net borrowings", NET_DEBT_NEW), ("Distributions declared, unpaid", DISTRIBUTION_NEW))):
        nb.write(5 + i, 0, label)
        for k in range(6):
            nb.write_number(5 + i, FIRST_COL - 1 + k, round(v0 * 0.97 ** k, 1), num)


# ---- the overlay (valuation) -----------------------------------------------------------------------------

def valuation(n: dict, vd: date, rate: float) -> dict:
    """One end of the range: the DCF at one discount rate."""
    ends = {k: e for k, e in enumerate(n["ends"])}
    df = [dcf.factors(ends, vd, rate, "end", "actual/actual")[k] for k in range(YEARS)]
    tv = n["fcf"][-1] * (1 + GROWTH) / (rate - GROWTH)
    flows = [f + (tv if k == YEARS - 1 else 0) for k, f in enumerate(n["fcf"])]
    pv_forecast = sum(n["fcf"][k] * df[k] for k in range(YEARS))
    pv_tv = tv * df[-1]
    ev = pv_forecast + pv_tv
    franking = sum(-n["tax"][k] * UTILISATION * df[k] for k in range(YEARS))
    cum = ev + franking - NET_DEBT
    return dict(rate=rate, df=df, tv=tv, flows=flows, ev=ev, pv_forecast=pv_forecast, pv_tv=pv_tv, franking=franking,
                cum=cum, equity=cum - DISTRIBUTION)


def write_overlay(wb, n: dict, fcf_ref, tax_ref, vd: date) -> dict:
    """Val_Inputs / DCF / Summary. fcf_ref / tax_ref(col_letter) -> formula text reading the client model."""
    b = wb.add_format({"bold": True})
    pct, num, dt = wb.add_format({"num_format": "0.00%"}), wb.add_format({"num_format": "#,##0.0"}), \
        wb.add_format({"num_format": "dd-mmm-yy"})
    four = wb.add_format({"num_format": "0.0000"})
    lo, hi = (valuation(n, vd, r) for r in RATES)
    vi = wb.add_worksheet("Val_Inputs")
    vi.write(0, 0, "Valuation assumptions", b)
    for c, h in ((2, "Low"), (3, "Mid"), (4, "High")):
        vi.write(2, c, h, b)
    vi.write(3, 0, "Valuation date"); vi.write_datetime(3, 2, vd, dt)
    vi.write(4, 0, "Discount rate (post-tax nominal WACC)")
    vi.write(4, 2, RATES[0], pct); vi.write_formula("D5", "=AVERAGE(C5,E5)", pct, sum(RATES) / 2); vi.write(4, 4, RATES[1], pct)
    vi.write(5, 0, "Terminal growth rate"); vi.write(5, 2, GROWTH, pct)
    vi.write(6, 0, "Net debt at valuation date"); vi.write(6, 1, "A$m"); vi.write(6, 2, NET_DEBT, num)
    vi.write(7, 0, "Declared distribution (unpaid at valuation date)"); vi.write(7, 1, "A$m"); vi.write(7, 2, DISTRIBUTION, num)
    vi.write(8, 0, "Franking credit utilisation (gamma)"); vi.write(8, 2, UTILISATION, pct)

    d = wb.add_worksheet("DCF")
    d.write(0, 0, "Discounted cash flow", b)
    d.write(2, 0, "Period ending", b)
    for k, e in enumerate(n["ends"]):
        d.write_datetime(2, FIRST_COL + k, e, dt)
    last = COL(FIRST_COL + YEARS - 1)
    d.write(4, 0, "Unlevered free cash flow (client model)"); d.write(4, 1, "A$m")
    d.write(5, 0, "Tax paid (client model)"); d.write(5, 1, "A$m")
    d.write(6, 0, "Franking credits utilised"); d.write(6, 1, "A$m")
    for k in range(YEARS):
        c = COL(FIRST_COL + k)
        d.write_formula(f"{c}5", fcf_ref(c), num, n["fcf"][k])
        d.write_formula(f"{c}6", tax_ref(c), num, n["tax"][k])
        d.write_formula(f"{c}7", f"=-{c}6*Val_Inputs!$C$9", num, -n["tax"][k] * UTILISATION)
    # one block per end of the range: terminal value, valuation cash flow, discount factor, present value
    for base, (v, rate_cell, end) in ((9, (lo, "Val_Inputs!$C$5", "low")), (16, (hi, "Val_Inputs!$E$5", "high"))):
        d.write(base - 1, 0, f"Valuation at the {end} end (discount rate {rate_cell.split('!')[1].replace('$', '')})", b)
        tv_r, vcf_r, df_r, pv_r = base + 1, base + 2, base + 3, base + 4  # Excel rows
        for r, lab in ((tv_r, "Terminal value"), (vcf_r, "Valuation cash flow"), (df_r, "Discount factor"),
                       (pv_r, "Present value")):
            d.write(r - 1, 0, lab)
            d.write(r - 1, 1, "" if r == df_r else "A$m")
        for k in range(YEARS):
            c = COL(FIRST_COL + k)
            if k == YEARS - 1:
                d.write_formula(f"{c}{tv_r}", f"={c}5*(1+Val_Inputs!$C$6)/({rate_cell}-Val_Inputs!$C$6)", num, v["tv"])
            else:
                d.write(f"{c}{tv_r}", 0, num)
            d.write_formula(f"{c}{vcf_r}", f"={c}5+{c}{tv_r}", num, v["flows"][k])
            d.write_formula(f"{c}{df_r}", f"=1/(1+{rate_cell})^YEARFRAC(Val_Inputs!$C$4,{c}$3,1)", four, v["df"][k])
            d.write_formula(f"{c}{pv_r}", f"={c}{vcf_r}*{c}{df_r}", num, v["flows"][k] * v["df"][k])
        v["rows"] = dict(tv=tv_r, vcf=vcf_r, df=df_r, pv=pv_r)

    s = wb.add_worksheet("Summary")
    s.write(0, 0, "Valuation summary (A$m)", b)
    for c, h in enumerate(["", "", "Low", "Mid", "High"]):
        s.write(2, c, h, b)
    L, H = lo["rows"], hi["rows"]
    items = [(4, "Enterprise value", f"=SUM(DCF!D{L['pv']}:{last}{L['pv']})", f"=SUM(DCF!D{H['pv']}:{last}{H['pv']})",
              lo["ev"], hi["ev"]),
             (5, "Value of franking credits", f"=SUMPRODUCT(DCF!D7:{last}7,DCF!D{L['df']}:{last}{L['df']})",
              f"=SUMPRODUCT(DCF!D7:{last}7,DCF!D{H['df']}:{last}{H['df']})", lo["franking"], hi["franking"]),
             (6, "Less: net debt", "=-Val_Inputs!$C$7", "=-Val_Inputs!$C$7", -NET_DEBT, -NET_DEBT),
             (7, "Equity value (cum-distribution)", "=SUM(C4:C6)", "=SUM(E4:E6)", lo["cum"], hi["cum"]),
             (8, "Less: declared distribution", "=-Val_Inputs!$C$8", "=-Val_Inputs!$C$8", -DISTRIBUTION, -DISTRIBUTION),
             (9, "Equity value (ex-distribution)", "=C7+C8", "=E7+E8", lo["equity"], hi["equity"])]
    for r, label, f_lo, f_hi, v_lo, v_hi in items:
        s.write(r - 1, 0, label)
        s.write(r - 1, 1, "A$m")
        s.write_formula(f"C{r}", f_lo, num, v_lo)
        s.write_formula(f"D{r}", f"=AVERAGE(C{r},E{r})", num, (v_lo + v_hi) / 2)
        s.write_formula(f"E{r}", f_hi, num, v_hi)
    s.write(10, 0, "The mid is the midpoint of the low and the high")
    return {"low": lo, "high": hi, "mid": {k: (lo[k] + hi[k]) / 2 for k in
                                           ("ev", "pv_forecast", "pv_tv", "tv", "franking", "cum", "equity")}}


def add_external_link(path: Path, target: str, sheets: list[str], cached: dict[tuple[str, str], float]) -> None:
    """Add xl/externalLinks/externalLink1.xml (with cached values) so [1]Sheet!A1 in formulas resolves,
    as Excel would have saved it. XlsxWriter writes the formulas but not the link part."""
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    by_sheet = {s: {} for s in sheets}
    for (s, addr), val in cached.items():
        by_sheet[s][addr] = val
    data = []
    for i, s in enumerate(sheets):
        rows = {}
        for addr, val in by_sheet[s].items():
            rows.setdefault(int(re.sub(r"\D", "", addr)), []).append((addr, val))
        body = "".join(f'<row r="{r}">' + "".join(f'<cell r="{a}"><v>{v!r}</v></cell>' for a, v in sorted(cells))
                       + "</row>" for r, cells in sorted(rows.items()))
        data.append(f'<sheetData sheetId="{i}">{body}</sheetData>')
    link = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<externalLink xmlns="{ns}" xmlns:r="{rns}">'
            f'<externalBook r:id="rId1"><sheetNames>' + "".join(f'<sheetName val="{s}"/>' for s in sheets)
            + f'</sheetNames><sheetDataSet>{"".join(data)}</sheetDataSet></externalBook></externalLink>')
    link_rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.'
                 'openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.'
                 'openxmlformats.org/officeDocument/2006/relationships/externalLinkPath" '
                 f'Target="{target}" TargetMode="External"/></Relationships>')
    src = zipfile.ZipFile(path)
    parts = {name: src.read(name) for name in src.namelist()}
    src.close()
    wb_xml = parts["xl/workbook.xml"].decode()
    wb_xml = wb_xml.replace("</sheets>", '</sheets><externalReferences><externalReference r:id="rIdExt1"/>'
                                         "</externalReferences>", 1)
    parts["xl/workbook.xml"] = wb_xml.encode()
    rels = parts["xl/_rels/workbook.xml.rels"].decode().replace(
        "</Relationships>", '<Relationship Id="rIdExt1" Type="http://schemas.openxmlformats.org/officeDocument/'
                            '2006/relationships/externalLink" Target="externalLinks/externalLink1.xml"/>'
                            "</Relationships>")
    parts["xl/_rels/workbook.xml.rels"] = rels.encode()
    ct = parts["[Content_Types].xml"].decode().replace(
        "</Types>", '<Override PartName="/xl/externalLinks/externalLink1.xml" ContentType="application/vnd.'
                    'openxmlformats-officedocument.spreadsheetml.externalLink+xml"/></Types>')
    parts["[Content_Types].xml"] = ct.encode()
    parts["xl/externalLinks/externalLink1.xml"] = link.encode()
    parts["xl/externalLinks/_rels/externalLink1.xml.rels"] = link_rels.encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, blob in parts.items():
            z.writestr(name, blob)


# ---- the report ------------------------------------------------------------------------------------------

def m(x: float) -> str:
    return f"{x:,.1f}"


def report_content(v: dict, n: dict) -> dict:
    lo, hi, mid = v["low"], v["high"], v["mid"]
    share = mid["franking"] / mid["equity"]
    return {
        "title": "Asset A Pty Ltd",
        "subtitle": f"Independent valuation as at {VD.day} {VD:%B %Y}  |  Project Alpha",
        "client": "Holdco A Pty Ltd",
        "letter": [
            "Private and confidential",
            "The Directors, Holdco A Pty Ltd",
            "Dear Directors",
            "You have asked us to assess the fair market value of 100% of the ordinary equity in Asset A Pty Ltd as at",
            "30 June 2025, on an ex-distribution basis, for the purposes of financial reporting. Our conclusion is set",
            "out in the executive summary, and this letter should be read with the full report.",
        ],
        "scope": [
            "Our work covered the Company's FY25 business plan model and the valuation assumptions in section 2.",
            "We relied on information provided by management and have not audited it. Values are in Australian",
            "dollars, millions (A$m), unless stated otherwise.",
        ],
        "summary": [
            "We have been engaged by Holdco A Pty Ltd to assess the fair market value of 100% of the ordinary",
            f"equity in Asset A Pty Ltd (the Company) as at {VD.day} {VD:%B %Y} (the Valuation Date).",
            f"We have assessed the equity value, on an ex-distribution basis, to be in the range of A${m(lo['equity'])}m",
            f"to A${m(hi['equity'])}m, with a midpoint of A${m(mid['equity'])}m. On a cum-distribution basis, including",
            f"the declared distribution of A${m(DISTRIBUTION)}m, the midpoint would be A${m(mid['cum'])}m.",
            f"The equity value includes the value of franking credits of A${m(mid['franking'])}m at the midpoint",
            f"({share * 100:.1f}% of the equity value).",
        ],
        "summary_table": [["A$m", "Low", "Mid", "High"],
                          ["Enterprise value", m(lo["ev"]), m(mid["ev"]), m(hi["ev"])],
                          ["Value of franking credits", m(lo["franking"]), m(mid["franking"]), m(hi["franking"])],
                          ["Less: net debt", f"({m(NET_DEBT)})", f"({m(NET_DEBT)})", f"({m(NET_DEBT)})"],
                          ["Less: declared distribution", f"({m(DISTRIBUTION)})", f"({m(DISTRIBUTION)})", f"({m(DISTRIBUTION)})"],
                          ["Equity value (ex-distribution)", m(lo["equity"]), m(mid["equity"]), m(hi["equity"])]],
        "method": [
            "Our primary valuation approach is a discounted cash flow (DCF) analysis of the unlevered free cash",
            "flows in the Company's FY25 business plan model, over the forecast period to FY45.",
            "A terminal value is calculated at the end of the forecast using the Gordon growth method,",
            "applied to the final forecast year's cash flow.",
            "Cash flows are discounted to the Valuation Date at a post-tax nominal WACC, end of period.",
            f"At the midpoint, the terminal value is A${m(mid['tv'])}m at 30 June 2045. The present value of the",
            f"forecast cash flows is A${m(mid['pv_forecast'])}m and the present value of the terminal value is",
            f"A${m(mid['pv_tv'])}m.",
        ],
        "assumptions": [["Assumption", "Value", "Basis"],
                        ["Discount rate", f"{RATES[1] * 100:.2f}% - {RATES[0] * 100:.2f}%", "Post-tax nominal WACC"],
                        ["Terminal growth rate", f"{GROWTH * 100:.2f}%", "Long-term CPI"],
                        ["Franking credit utilisation", f"{UTILISATION * 100:.0f}%", "Gamma"],
                        ["Forecast period", "FY26 - FY45", "20 years"],
                        ["Corporate tax rate", "30.00%", "Statutory"]],
    }


def table_png(rows: list[list[str]], width: float = 7.0) -> bytes:
    """A table as a picture, the way a report pastes one in: no text layer, so it has to be read from the image."""
    plt.rcParams["text.parse_math"] = False
    fig = plt.figure(figsize=(width, 0.42 * len(rows) + 0.3), dpi=200)
    ax = fig.add_axes([0, 0, 1, 1]); ax.axis("off")
    t = ax.table(cellText=rows[1:], colLabels=rows[0], loc="center", cellLoc="left")
    t.auto_set_font_size(False); t.set_fontsize(10); t.scale(1, 1.5)
    for (r, _), cell in t.get_celld().items():
        cell.set_edgecolor("#c4c4cd")
        if r == 0:
            cell.set_facecolor("#2e2e38"); cell.get_text().set_color("white")
    buf = io.BytesIO(); fig.savefig(buf, format="png"); plt.close(fig)
    return buf.getvalue()


def write_pdf(path: Path, c: dict) -> None:
    plt.rcParams["pdf.fonttype"] = 42  # TrueType, so the text layer can be read back
    plt.rcParams["text.parse_math"] = False  # "A$2,135.4m to A$2,475.9m" is not maths
    with PdfPages(path) as pdf:
        def page(title, n):
            fig = plt.figure(figsize=(8.27, 11.69))
            fig.text(0.08, 0.93, title, fontsize=17, weight="bold")
            fig.text(0.5, 0.03, f"Page {n}", fontsize=8, ha="center", color="#747480")
            return fig

        def text(fig, lines, y, size=10.5):
            for ln in lines:
                fig.text(0.08, y, ln, fontsize=size); y -= 0.022
            return y

        def table(fig, rows, y, h):
            ax = fig.add_axes([0.08, y - h, 0.84, h]); ax.axis("off")
            t = ax.table(cellText=rows[1:], colLabels=rows[0], loc="upper left", cellLoc="right")
            t.auto_set_font_size(False); t.set_fontsize(9.5); t.scale(1, 1.4)
            for (r, col), cell in t.get_celld().items():
                cell.set_edgecolor("#c4c4cd")
                if col == 0:
                    cell._loc = "left"

        fig = plt.figure(figsize=(8.27, 11.69))
        fig.text(0.08, 0.6, c["title"], fontsize=24, weight="bold")
        fig.text(0.08, 0.56, c["subtitle"], fontsize=13)
        fig.text(0.08, 0.52, f"Prepared for {c['client']}  |  Final report", fontsize=10, color="#747480")
        fig.text(0.08, 0.1, "Synthetic test document: every name and number in it is fictional.", fontsize=8)
        pdf.savefig(fig); plt.close(fig)

        # the transmittal letter, with the scope of the work: context for every figure after it
        fig = page("Letter of transmittal", 2)
        y = text(fig, c["letter"], 0.88)
        fig.text(0.08, y - 0.02, "Scope of our work", fontsize=10.5, weight="bold")
        y = text(fig, c["scope"], y - 0.05)
        text(fig, ["Yours faithfully", "The engagement team"], y - 0.03)
        pdf.savefig(fig); plt.close(fig)

        fig = page("1. Executive summary", 3)
        y = text(fig, c["summary"], 0.88)
        fig.text(0.08, y - 0.02, "Table 1: Valuation summary", fontsize=10, weight="bold")
        table(fig, c["summary_table"], y - 0.035, 0.2)
        pdf.savefig(fig); plt.close(fig)

        fig = page("2. Valuation approach", 4)
        y = text(fig, c["method"], 0.88)
        fig.text(0.08, y - 0.02, "Table 2: Key valuation assumptions", fontsize=10, weight="bold")
        # pasted as a picture: the discount rate is only here, so it has to be read from the image
        from PIL import Image
        img = Image.open(io.BytesIO(table_png(c["assumptions"])))
        ax = fig.add_axes([0.08, y - 0.3, 0.84, 0.26]); ax.axis("off"); ax.imshow(img)
        pdf.savefig(fig); plt.close(fig)


def write_pptx(path: Path, c: dict) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = c["title"]
    s.placeholders[1].text = c["subtitle"]

    def slide(title):
        sl = prs.slides.add_slide(prs.slide_layouts[5])
        sl.shapes.title.text = title
        return sl

    def table(sl, rows, top):
        t = sl.shapes.add_table(len(rows), len(rows[0]), Inches(0.6), Inches(top), Inches(9), Inches(0.4 * len(rows))).table
        for r, row in enumerate(rows):
            for k, val in enumerate(row):
                t.cell(r, k).text = val

    sl = slide("Letter of transmittal")
    tb = sl.shapes.add_textbox(Inches(0.6), Inches(1.3), Inches(12), Inches(4)).text_frame
    tb.word_wrap = True
    for i, ln in enumerate(c["letter"] + ["Scope of our work"] + c["scope"] + ["Yours faithfully", "The engagement team"]):
        p = tb.paragraphs[0] if i == 0 else tb.add_paragraph()
        p.text, p.font.size = ln, Pt(12)

    sl = slide("Executive summary")
    tb = sl.shapes.add_textbox(Inches(0.6), Inches(1.3), Inches(12), Inches(1.8)).text_frame
    tb.word_wrap = True
    tb.text = " ".join(c["summary"])
    for p in tb.paragraphs:
        p.font.size = Pt(13)
    table(sl, c["summary_table"], 3.6)

    sl = slide("Valuation approach and key assumptions")
    tb = sl.shapes.add_textbox(Inches(0.6), Inches(1.3), Inches(12), Inches(1.8)).text_frame
    tb.word_wrap = True
    for i, ln in enumerate(c["method"]):
        p = tb.paragraphs[0] if i == 0 else tb.add_paragraph()
        p.text, p.font.size = ln, Pt(12)
    table(sl, c["assumptions"], 4.2)
    prs.save(path)


# ---- write everything ------------------------------------------------------------------------------------

def inputs_rows(vd: date, volume0, growth, tariff0, cpi, opex_pct, capex, major, insurance=None):
    rows = [("Valuation date", vd, "date"), ("Opening volume", volume0, "m units"), ("Volume growth", growth, "%"),
            ("Tariff at start of forecast", tariff0, "A$"), ("CPI", cpi, "%"),
            ("Operating costs (% of revenue)", opex_pct, "%"), ("Maintenance capex", capex, "A$m"),
            ("Major maintenance (every 5 years)", major, "A$m"), ("Tax rate", 0.30, "%")]
    if insurance is not None:
        rows.append(("Insurance premium", insurance, "A$m"))
    return {"rows": rows}


def main() -> dict:
    OUT.mkdir(parents=True, exist_ok=True)
    prior_in = dict(volume0=42.0, growth=0.02, tariff0=6.50, cpi=0.025, opex_pct=0.18, capex=35.0, major=120.0)
    cur_in = dict(volume0=43.1, growth=0.021, tariff0=6.70, cpi=0.03, opex_pct=0.18, capex=38.0, major=120.0)
    prior = client_numbers(2026, **prior_in, insurance=None)
    current = client_numbers(2027, **cur_in, insurance=4.0)

    files = {}
    p = OUT / "AssetA_BP25_client_model_Jun25.xlsx"
    wb = xlsxwriter.Workbook(p)
    rows = write_client(wb, prior, inputs_rows(VD, **prior_in), insurance=False, balances=(NET_DEBT, DISTRIBUTION))
    wb.close(); files["prior client model"] = p

    p = OUT / "AssetA_BP26_client_model_Jun26.xlsx"
    wb = xlsxwriter.Workbook(p)
    write_client(wb, current, inputs_rows(VD_NEW, **cur_in, insurance=4.0), insurance=True,
                 balances=(NET_DEBT_NEW, DISTRIBUTION_NEW))
    wb.close(); files["current client model"] = p

    p = OUT / "AssetA_FY26_plan_rebuilt.xlsx"  # this year's model, rebuilt from the ground up
    wb = xlsxwriter.Workbook(p)
    write_rebuilt(wb, current, prior, VD_NEW)
    wb.close(); files["current client model (rebuilt)"] = p

    fcf_row, tax_row = rows["cf_fcf"][1], rows["cf_tax"][1]
    p = OUT / "Alpha_valuation_overlay_FY25.xlsx"
    wb = xlsxwriter.Workbook(p)
    v = write_overlay(wb, prior, lambda c: f"=[1]CashFlow!{c}{fcf_row}", lambda c: f"=[1]CashFlow!{c}{tax_row}", VD)
    wb.close()
    cached = {("CashFlow", f"{COL(FIRST_COL + k)}{r}"): prior[key][k] for k in range(YEARS)
              for r, key in ((fcf_row, "fcf"), (tax_row, "tax"))}
    add_external_link(p, "AssetA_BP25_client_model_Jun25.xlsx", ["Inputs", "Operations", "CashFlow"], cached)
    files["prior overlay (standalone)"] = p

    p = OUT / "AssetA_BP25_with_overlay.xlsx"
    wb = xlsxwriter.Workbook(p)
    write_client(wb, prior, inputs_rows(VD, **prior_in), insurance=False, balances=(NET_DEBT, DISTRIBUTION))
    write_overlay(wb, prior, lambda c: f"=CashFlow!{c}{fcf_row}", lambda c: f"=CashFlow!{c}{tax_row}", VD)
    wb.close(); files["prior client model with overlay inside"] = p

    content = report_content(v, prior)
    p = OUT / "AssetA_valuation_report_FY25.pdf"
    write_pdf(p, content); files["prior report (PDF)"] = p
    p = OUT / "AssetA_valuation_report_FY25.pptx"
    write_pptx(p, content); files["prior report (PPTX)"] = p

    for role, f in files.items():
        print(f"{role:40s} {f.relative_to(ROOT)}")
    lo, hi, mid = v["low"], v["high"], v["mid"]
    print(f"\nprior: equity (ex-distribution) A${m(lo['equity'])}m - A${m(hi['equity'])}m, mid A${m(mid['equity'])}m; "
          f"rates {RATES[1]:.2%} - {RATES[0]:.2%}, TGR {GROWTH:.2%}")
    return {"files": files, "valuation": v, "prior": prior, "current": current}


if __name__ == "__main__":
    main()
